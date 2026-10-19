"""Bounded, typed result export without Qt or database dependencies.

The caller captures typed rows on the appropriate owner thread and passes an
immutable snapshot or a lazy sequence of ``ManualQueryPage`` objects here.
Format contracts are versioned and preserve values that common interchange
formats otherwise coerce or lose.
"""

from __future__ import annotations

import base64
import csv
import datetime as _datetime
import json
import os
import tempfile
from dataclasses import dataclass
from decimal import Decimal
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator, Sequence
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from .manual_session import ManualColumn, ManualQueryPage


class ExportScope(str, Enum):
    SELECTED = "selected"
    CURRENT_PAGE = "current_page"
    FULL_QUERY = "full_query"


class ExportError(Exception):
    """A safe, machine-readable export failure; never includes cell contents."""

    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


@dataclass(frozen=True, slots=True)
class ExportTarget:
    query_id: str
    tab_id: str
    generation: int

    def __post_init__(self) -> None:
        if not isinstance(self.query_id, str) or not self.query_id or "\x00" in self.query_id:
            raise ExportError("INVALID_TARGET")
        if not isinstance(self.tab_id, str) or not self.tab_id or "\x00" in self.tab_id:
            raise ExportError("INVALID_TARGET")
        try:
            too_long = len(self.query_id.encode("utf-8")) > 4_096 or len(self.tab_id.encode("utf-8")) > 4_096
        except UnicodeError as error:
            raise ExportError("INVALID_TARGET") from error
        if too_long:
            raise ExportError("INVALID_TARGET")
        if isinstance(self.generation, bool) or not isinstance(self.generation, int) or self.generation < 0:
            raise ExportError("INVALID_TARGET")


@dataclass(frozen=True, slots=True)
class ExportSnapshot:
    """Immutable typed selection/current-page snapshot captured by the caller."""

    target: ExportTarget
    columns: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    page_number: int | None = None

    def __post_init__(self) -> None:
        if not isinstance(self.target, ExportTarget):
            raise ExportError("INVALID_TARGET")
        try:
            columns = tuple(self.columns)
            rows = tuple(tuple(row) for row in self.rows)
        except Exception as error:
            raise ExportError("INVALID_SNAPSHOT") from error
        if any(not isinstance(name, str) for name in columns):
            raise ExportError("INVALID_COLUMNS")
        if self.page_number is not None and (
            isinstance(self.page_number, bool)
            or not isinstance(self.page_number, int)
            or self.page_number < 1
        ):
            raise ExportError("INVALID_PAGE")
        object.__setattr__(self, "columns", columns)
        object.__setattr__(self, "rows", rows)


@dataclass(frozen=True, slots=True)
class ExportLimits:
    max_rows: int = 1_000_000
    max_pages: int = 100_000
    max_columns: int = 16_384
    max_snapshot_bytes: int = 50 * 1024 * 1024
    max_page_bytes: int = 50 * 1024 * 1024
    max_output_bytes: int = 512 * 1024 * 1024
    max_cell_bytes: int = 16 * 1024 * 1024
    max_xlsx_metadata_rows: int = 1_048_576

    def __post_init__(self) -> None:
        values = (
            self.max_rows,
            self.max_pages,
            self.max_columns,
            self.max_snapshot_bytes,
            self.max_page_bytes,
            self.max_output_bytes,
            self.max_cell_bytes,
            self.max_xlsx_metadata_rows,
        )
        if any(isinstance(value, bool) or not isinstance(value, int) or value <= 0 for value in values):
            raise ExportError("INVALID_LIMITS")


DEFAULT_LIMITS = ExportLimits()


@dataclass(frozen=True, slots=True)
class ExportReceipt:
    path: Path
    format: str
    scope: ExportScope
    rows: int
    bytes_written: int
    contract_version: int = 1


_CSV_MARKER = "~PTX-CELL/1~"
CSV_CONTRACT = "PengTools CSV typed-cell marker v1"
JSON_CONTRACT = "pengtools.typed-result-export"
XLSX_METADATA_SHEET = "_PT_EXPORT_V1"
_XLSX_META_SHEET = XLSX_METADATA_SHEET
_XLSX_CELL_LIMIT = 32_767
_XLSX_META_CHUNK_CHARS = 30_000
_XLSX_MAX_WORKSHEET_ROWS = 1_048_576
_SUPPORTED_FORMATS = frozenset({"csv", "json", "xlsx"})


def write_result_export(
    destination: str | os.PathLike[str],
    format: str,
    scope: ExportScope | str,
    *,
    target: ExportTarget,
    snapshot: ExportSnapshot | None = None,
    pages: Iterable[ManualQueryPage] | None = None,
    cancel_check: Callable[[], bool] | None = None,
    overwrite: bool = False,
    limits: ExportLimits = DEFAULT_LIMITS,
) -> ExportReceipt:
    """Stream a typed snapshot or complete page chain to an atomically published file.

    ``full_query`` requires page 1 through its explicit terminal page.  The
    function does no connection, driver, Qt, or active-session work.
    """

    selected_format = format.lower() if isinstance(format, str) else ""
    try:
        selected_scope = scope if isinstance(scope, ExportScope) else ExportScope(scope)
    except (ValueError, TypeError) as error:
        raise ExportError("INVALID_SCOPE") from error
    if selected_format not in _SUPPORTED_FORMATS:
        raise ExportError("UNSUPPORTED_FORMAT")
    if not isinstance(target, ExportTarget) or not isinstance(limits, ExportLimits):
        raise ExportError("INVALID_REQUEST")
    if not isinstance(overwrite, bool):
        raise ExportError("INVALID_REQUEST")
    if selected_scope is ExportScope.FULL_QUERY:
        if snapshot is not None or pages is None:
            raise ExportError("INVALID_SOURCE")
    elif snapshot is None or pages is not None:
        raise ExportError("INVALID_SOURCE")
    if cancel_check is not None and not callable(cancel_check):
        raise ExportError("INVALID_CANCELLATION")

    destination_path = Path(destination)
    parent = destination_path.parent
    if not parent.is_dir() or destination_path.name in {"", ".", ".."}:
        raise ExportError("DESTINATION_UNAVAILABLE")
    if not overwrite and os.path.lexists(destination_path):
        raise ExportError("DESTINATION_EXISTS")

    _check_cancel(cancel_check)
    temp_path: Path | None = None
    rows_written = 0
    output_bytes = 0
    try:
        handle = tempfile.NamedTemporaryFile(
            mode="wb", prefix=".pengtools-export-", suffix=".tmp", dir=parent, delete=False
        )
        temp_path = Path(handle.name)
        handle.close()

        if selected_scope is ExportScope.FULL_QUERY:
            assert pages is not None
            output_bytes, rows_written, _page_count = _write_full_query(
                temp_path,
                selected_format,
                pages,
                target,
                cancel_check,
                limits,
            )
        else:
            assert snapshot is not None
            output_bytes, rows_written = _write_snapshot(
                temp_path,
                selected_format,
                selected_scope,
                snapshot,
                target,
                cancel_check,
                limits,
            )

        _check_cancel(cancel_check)
        _fsync_file(temp_path)
        size = temp_path.stat().st_size
        if size > limits.max_output_bytes:
            raise ExportError("OUTPUT_LIMIT")
        _check_cancel(cancel_check)
        if overwrite:
            os.replace(temp_path, destination_path)
            temp_path = None
        else:
            try:
                # Same-directory hard-link creation is atomic and fails if the
                # destination appeared after the early existence check.
                os.link(temp_path, destination_path)
            except FileExistsError as error:
                raise ExportError("DESTINATION_EXISTS") from error
            except OSError as error:
                raise ExportError("ATOMIC_CREATE_UNAVAILABLE") from error
            try:
                temp_path.unlink()
            except OSError:
                # Publication succeeded; the finally block gets one more
                # chance to remove only this operation's temporary name.
                pass
        _fsync_directory(parent)
        return ExportReceipt(
            path=destination_path,
            format=selected_format,
            scope=selected_scope,
            rows=rows_written,
            bytes_written=size,
        )
    except ExportError:
        raise
    except Exception as error:
        # Errors from codecs, iterable sources, or writers can include raw
        # values in provider-specific text; expose only a stable safe code.
        raise ExportError("EXPORT_FAILED") from error
    finally:
        if temp_path is not None:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass


def encode_typed_cell(value: Any) -> dict[str, Any]:
    """Return the v1 typed-cell JSON object used by JSON/XLSX contracts."""

    if value is None:
        return {"t": "null"}
    if isinstance(value, bool):
        return {"t": "bool", "v": value}
    if isinstance(value, int):
        return {"t": "int", "v": str(value)}
    if isinstance(value, Decimal):
        return {"t": "decimal", "v": str(value)}
    if isinstance(value, float):
        return {"t": "float", "v": value.hex()}
    if isinstance(value, str):
        return {"t": "str", "v": value}
    if isinstance(value, bytes):
        return {"t": "bytes", "v": base64.b64encode(value).decode("ascii")}
    if isinstance(value, _datetime.datetime):
        return {
            "t": "datetime",
            "v": value.isoformat(),
            "fold": value.fold,
            "tz": _timezone_contract(value),
        }
    if isinstance(value, _datetime.date):
        return {"t": "date", "v": value.isoformat()}
    if isinstance(value, _datetime.time):
        return {
            "t": "time",
            "v": value.isoformat(),
            "fold": value.fold,
            "tz": _timezone_contract(value),
        }
    raise ExportError("UNSUPPORTED_VALUE")


def decode_typed_cell(cell: dict[str, Any]) -> Any:
    """Decode a v1 typed-cell object, useful to consumers of the export contract."""

    try:
        kind = cell["t"]
        value = cell.get("v")
        if kind == "null":
            return None
        if kind == "bool" and isinstance(value, bool):
            return value
        if kind == "int" and isinstance(value, str):
            return int(value)
        if kind == "decimal" and isinstance(value, str):
            return Decimal(value)
        if kind == "float" and isinstance(value, str):
            return float.fromhex(value)
        if kind == "str" and isinstance(value, str):
            return value
        if kind == "bytes" and isinstance(value, str):
            return base64.b64decode(value, validate=True)
        if kind == "date" and isinstance(value, str):
            return _datetime.date.fromisoformat(value)
        if kind in {"datetime", "time"} and isinstance(value, str):
            parsed = _datetime.datetime.fromisoformat(value) if kind == "datetime" else _datetime.time.fromisoformat(value)
            parsed = parsed.replace(fold=int(cell.get("fold", 0)))
            tz = cell.get("tz")
            if isinstance(tz, dict) and tz.get("kind") == "zoneinfo":
                parsed = parsed.replace(tzinfo=ZoneInfo(tz["key"]))
                offset = parsed.utcoffset()
                if offset is None or int(offset.total_seconds()) != int(tz["offset_seconds"]):
                    raise ValueError("timezone offset mismatch")
            elif isinstance(tz, dict) and tz.get("kind") == "offset":
                offset = _datetime.timedelta(seconds=int(tz["seconds"]))
                name = tz.get("name")
                fixed = _datetime.timezone(offset, name) if isinstance(name, str) else _datetime.timezone(offset)
                parsed = parsed.replace(tzinfo=fixed)
            return parsed
    except (ArithmeticError, KeyError, TypeError, ValueError, ZoneInfoNotFoundError) as error:
        raise ExportError("INVALID_TYPED_CELL") from error
    raise ExportError("INVALID_TYPED_CELL")


def decode_csv_cell(value: str) -> Any:
    """Decode one v1 CSV data cell; ordinary strings are returned unchanged."""

    if not value.startswith(_CSV_MARKER):
        return value
    encoded = value[len(_CSV_MARKER) :]
    try:
        payload = base64.b64decode(encoded + "=" * (-len(encoded) % 4), altchars=b"-_", validate=True)
        cell = json.loads(payload.decode("utf-8"))
    except Exception as error:
        raise ExportError("INVALID_CSV_CELL") from error
    if not isinstance(cell, dict):
        raise ExportError("INVALID_CSV_CELL")
    return decode_typed_cell(cell)


def _timezone_contract(value: _datetime.datetime | _datetime.time) -> dict[str, Any]:
    tz = value.tzinfo
    if tz is None:
        return {"kind": "naive"}
    offset = value.utcoffset()
    if offset is None:
        raise ExportError("UNSUPPORTED_TIMEZONE")
    if isinstance(tz, ZoneInfo):
        return {"kind": "zoneinfo", "key": tz.key, "offset_seconds": int(offset.total_seconds())}
    if not isinstance(tz, _datetime.timezone):
        raise ExportError("UNSUPPORTED_TIMEZONE")
    seconds = int(offset.total_seconds())
    if offset.total_seconds() != seconds:
        raise ExportError("UNSUPPORTED_TIMEZONE")
    return {"kind": "offset", "seconds": seconds, "name": value.tzname()}


def _check_cancel(cancel_check: Callable[[], bool] | None) -> None:
    if cancel_check is None:
        return
    try:
        cancelled = cancel_check()
    except Exception as error:
        raise ExportError("CANCELLATION_CHECK_FAILED") from error
    if cancelled:
        raise ExportError("CANCELLED")


def _write_snapshot(
    temp_path: Path,
    selected_format: str,
    scope: ExportScope,
    snapshot: ExportSnapshot,
    target: ExportTarget,
    cancel_check: Callable[[], bool] | None,
    limits: ExportLimits,
) -> tuple[int, int]:
    if snapshot.target != target:
        raise ExportError("TARGET_MISMATCH")
    if scope is ExportScope.CURRENT_PAGE and snapshot.page_number is None:
        raise ExportError("INVALID_PAGE")
    if scope is ExportScope.SELECTED and not snapshot.rows:
        raise ExportError("EMPTY_SELECTION")
    names = _validate_column_names(snapshot.columns, limits)
    if len(snapshot.rows) > limits.max_rows:
        raise ExportError("ROW_LIMIT")
    source_bytes = 0
    for row in snapshot.rows:
        _check_cancel(cancel_check)
        source_bytes += _validate_row(row, len(names), limits)
        if source_bytes > limits.max_snapshot_bytes:
            raise ExportError("SNAPSHOT_LIMIT")
    _check_cancel(cancel_check)
    output_bytes, rows_written = _write_rows(
        temp_path, selected_format, scope, target, names, iter(snapshot.rows), cancel_check, limits
    )
    return output_bytes, rows_written


def _write_full_query(
    temp_path: Path,
    selected_format: str,
    pages: Iterable[ManualQueryPage],
    target: ExportTarget,
    cancel_check: Callable[[], bool] | None,
    limits: ExportLimits,
) -> tuple[int, int, int]:
    try:
        iterator = iter(pages)
        first = next(iterator)
    except StopIteration as error:
        raise ExportError("INCOMPLETE_PAGE_CHAIN") from error
    except ExportError:
        raise
    except Exception as error:
        raise ExportError("PAGE_SOURCE_FAILED") from error

    if not isinstance(first, ManualQueryPage):
        raise ExportError("INVALID_PAGE")
    names = _validate_page(first, target, expected_page=1, prior_columns=None, limits=limits)
    first_columns = tuple(first.columns)
    rows_written = 0
    page_count = 0

    def all_rows() -> Iterator[tuple[Any, ...]]:
        nonlocal rows_written, page_count
        page = first
        expected_page = 1
        while True:
            _check_cancel(cancel_check)
            page_names = _validate_page(page, target, expected_page, first_columns, limits)
            if page_names != names:
                raise ExportError("COLUMN_MISMATCH")
            page_count += 1
            if page_count > limits.max_pages:
                raise ExportError("PAGE_LIMIT")
            page_bytes = 0
            for row in page.rows:
                _check_cancel(cancel_check)
                page_bytes += _validate_row(row, len(names), limits)
                if page_bytes > limits.max_page_bytes:
                    raise ExportError("PAGE_SIZE_LIMIT")
                rows_written += 1
                if rows_written > limits.max_rows:
                    raise ExportError("ROW_LIMIT")
                yield row
            if not page.has_more:
                try:
                    next(iterator)
                except StopIteration:
                    return
                except Exception as error:
                    raise ExportError("PAGE_SOURCE_FAILED") from error
                raise ExportError("INCOMPLETE_PAGE_CHAIN")
            expected_page += 1
            try:
                page = next(iterator)
            except StopIteration as error:
                raise ExportError("INCOMPLETE_PAGE_CHAIN") from error
            except Exception as error:
                raise ExportError("PAGE_SOURCE_FAILED") from error

    output_bytes, _ = _write_rows(
        temp_path,
        selected_format,
        ExportScope.FULL_QUERY,
        target,
        names,
        all_rows(),
        cancel_check,
        limits,
    )
    return output_bytes, rows_written, page_count


def _validate_page(
    page: ManualQueryPage,
    target: ExportTarget,
    expected_page: int,
    prior_columns: tuple[ManualColumn, ...] | None,
    limits: ExportLimits,
) -> tuple[str, ...]:
    if not isinstance(page, ManualQueryPage):
        raise ExportError("INVALID_PAGE")
    if page.query_id != target.query_id or page.tab_id != target.tab_id or page.generation != target.generation:
        raise ExportError("TARGET_MISMATCH")
    if isinstance(page.page_number, bool) or not isinstance(page.page_number, int) or page.page_number != expected_page:
        raise ExportError("INCOMPLETE_PAGE_CHAIN")
    if not isinstance(page.truncated_by_budget, bool):
        raise ExportError("INVALID_PAGE")
    if page.truncated_by_budget:
        raise ExportError("SOURCE_TRUNCATED")
    if not isinstance(page.has_more, bool):
        raise ExportError("INVALID_PAGE")
    if (
        isinstance(page.page_size, bool)
        or not isinstance(page.page_size, int)
        or page.page_size < 1
        or isinstance(page.result_bytes, bool)
        or not isinstance(page.result_bytes, int)
        or page.result_bytes < 0
    ):
        raise ExportError("INVALID_PAGE")
    if not isinstance(page.columns, tuple) or not isinstance(page.rows, tuple):
        raise ExportError("INVALID_PAGE")
    if any(not isinstance(row, tuple) for row in page.rows) or len(page.rows) > page.page_size:
        raise ExportError("INVALID_PAGE")
    if any(not isinstance(column, ManualColumn) for column in page.columns):
        raise ExportError("INVALID_COLUMNS")
    names = _validate_column_names(tuple(column.name for column in page.columns), limits)
    if prior_columns is not None:
        try:
            metadata_matches = tuple(page.columns) == prior_columns
        except Exception as error:
            raise ExportError("COLUMN_METADATA_MISMATCH") from error
        if not metadata_matches:
            raise ExportError("COLUMN_METADATA_MISMATCH")
    return names


def _validate_column_names(names: Sequence[str], limits: ExportLimits) -> tuple[str, ...]:
    result = tuple(names)
    if len(result) > limits.max_columns:
        raise ExportError("COLUMN_LIMIT")
    for name in result:
        if not isinstance(name, str):
            raise ExportError("INVALID_COLUMNS")
        if _utf16_units(name) > _XLSX_CELL_LIMIT:
            raise ExportError("COLUMN_NAME_TOO_LONG")
        _check_cell_size(name, limits)
    return result


def _validate_row(row: Sequence[Any], column_count: int, limits: ExportLimits) -> int:
    if len(row) != column_count:
        raise ExportError("ROW_SHAPE_MISMATCH")
    total = 0
    for value in row:
        _check_cell_size(value, limits)
        cell = encode_typed_cell(value)
        total += len(_json_bytes(cell))
        if total > limits.max_page_bytes:
            raise ExportError("PAGE_SIZE_LIMIT")
    return total


def _check_cell_size(value: Any, limits: ExportLimits) -> None:
    if isinstance(value, str):
        size = len(value.encode("utf-8"))
    elif isinstance(value, bytes):
        size = len(value)
    elif isinstance(value, Decimal):
        size = len(str(value).encode("ascii", errors="strict"))
    elif isinstance(value, (int, float, bool)) or value is None:
        size = len(_json_bytes(encode_typed_cell(value)))
    elif isinstance(value, (_datetime.date, _datetime.time, _datetime.datetime)):
        size = len(_json_bytes(encode_typed_cell(value)))
    else:
        raise ExportError("UNSUPPORTED_VALUE")
    if size > limits.max_cell_bytes:
        raise ExportError("CELL_LIMIT")


def _write_rows(
    temp_path: Path,
    selected_format: str,
    scope: ExportScope,
    target: ExportTarget,
    columns: tuple[str, ...],
    rows: Iterable[tuple[Any, ...]],
    cancel_check: Callable[[], bool] | None,
    limits: ExportLimits,
) -> tuple[int, int]:
    if selected_format == "xlsx":
        count = _write_xlsx(temp_path, scope, target, columns, rows, cancel_check, limits)
        return temp_path.stat().st_size, count

    with temp_path.open("wb") as binary:
        writer = _Utf8BudgetWriter(binary, limits.max_output_bytes)
        if selected_format == "csv":
            csv_writer = csv.writer(writer, lineterminator="\r\n")
            csv_writer.writerow(tuple(_csv_encode_cell(name) for name in columns))
            count = 0
            for row in rows:
                _check_cancel(cancel_check)
                csv_writer.writerow(tuple(_csv_encode_cell(value) for value in row))
                count += 1
            writer.flush()
            return writer.bytes_written, count

        _write_json_header(writer, scope, target, columns)
        count = 0
        first_row = True
        for row in rows:
            _check_cancel(cancel_check)
            if not first_row:
                writer.write(",")
            writer.write(_json_text([encode_typed_cell(value) for value in row]))
            first_row = False
            count += 1
        writer.write("]}")
        writer.flush()
        return writer.bytes_written, count


def _csv_encode_cell(value: Any) -> str:
    if isinstance(value, str) and not value.startswith(_CSV_MARKER) and not _is_formula_like(value):
        return value
    payload = _json_bytes(encode_typed_cell(value))
    token = base64.urlsafe_b64encode(payload).decode("ascii").rstrip("=")
    return _CSV_MARKER + token


def _is_formula_like(value: str) -> bool:
    candidate = value.lstrip(" \t\r\n")
    return bool(candidate) and candidate[0] in "=+-@"


def _write_json_header(
    writer: _Utf8BudgetWriter,
    scope: ExportScope,
    target: ExportTarget,
    columns: tuple[str, ...],
) -> None:
    header = {
        "contract": JSON_CONTRACT,
        "version": 1,
        "scope": scope.value,
        "target": {"query_id": target.query_id, "tab_id": target.tab_id, "generation": target.generation},
        "columns": columns,
        "rows": None,
    }
    prefix = _json_text(header)
    if not prefix.endswith("null}"):
        raise ExportError("EXPORT_FAILED")
    writer.write(prefix[:-5] + "[")


def _write_xlsx(
    temp_path: Path,
    scope: ExportScope,
    target: ExportTarget,
    columns: tuple[str, ...],
    rows: Iterable[tuple[Any, ...]],
    cancel_check: Callable[[], bool] | None,
    limits: ExportLimits,
) -> int:
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
    except Exception as error:
        raise ExportError("XLSX_UNAVAILABLE") from error

    if len(columns) > 16_384:
        raise ExportError("XLSX_COLUMN_LIMIT")
    metadata_row_limit = min(limits.max_xlsx_metadata_rows, _XLSX_MAX_WORKSHEET_ROWS)
    if metadata_row_limit < 2:
        raise ExportError("XLSX_METADATA_LIMIT")
    data_row_limit = min(limits.max_rows, _XLSX_MAX_WORKSHEET_ROWS - 1)
    header_bytes = sum(len(name.encode("utf-8")) for name in columns)
    header_bytes += len(target.query_id.encode("utf-8")) + len(target.tab_id.encode("utf-8"))
    expanded_bytes = 12 * header_bytes + 1_024
    if expanded_bytes > limits.max_output_bytes:
        raise ExportError("OUTPUT_LIMIT")
    workbook = Workbook(write_only=True)
    saved = False
    try:
        data_sheet = workbook.create_sheet("Data")
        metadata_sheet = workbook.create_sheet(_XLSX_META_SHEET)
        metadata_sheet.sheet_state = "veryHidden"

        def safe_string(sheet: Any, value: str) -> Any:
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = "s"
            return cell

        data_sheet.append([safe_string(data_sheet, name) for name in columns])
        metadata_sheet.append(
            [
                "PengTools typed result export",
                1,
                scope.value,
                safe_string(metadata_sheet, target.query_id),
                safe_string(metadata_sheet, target.tab_id),
                target.generation,
            ]
        )
        metadata_sheet.append(["record_index", "column_index", "typed_cell_json", "chunk_index", "chunk_count"])

        count = 0
        metadata_rows = 2
        for row in rows:
            _check_cancel(cancel_check)
            if count >= data_row_limit:
                raise ExportError("XLSX_ROW_LIMIT")
            visible: list[Any] = []
            for column_index, value in enumerate(row):
                typed = encode_typed_cell(value)
                typed_json = _json_text(typed)
                # Write-only XLSX still spools XML to disk. Bound the two payload
                # copies (visible cell and typed sidecar) plus worst-case XML
                # escaping before asking the workbook writer to append them.
                expanded_bytes += 12 * len(typed_json.encode("utf-8")) + 80
                if expanded_bytes > limits.max_output_bytes:
                    raise ExportError("OUTPUT_LIMIT")
                visible.append(_xlsx_visible_cell(data_sheet, value))
                chunks = _split_xlsx_chunks(typed_json)
                chunk_count = len(chunks)
                if metadata_rows + chunk_count > metadata_row_limit:
                    raise ExportError("XLSX_METADATA_LIMIT")
                for chunk_index, chunk in enumerate(chunks):
                    metadata_sheet.append([count, column_index, safe_string(metadata_sheet, chunk), chunk_index, chunk_count])
                    metadata_rows += 1
            data_sheet.append(visible)
            count += 1

        _check_cancel(cancel_check)
        workbook.save(temp_path)
        saved = True
    except Exception as error:
        if isinstance(error, ExportError):
            raise
        raise ExportError("XLSX_WRITE_FAILED") from error
    finally:
        if not saved:
            _cleanup_xlsx_spools(workbook)
    if temp_path.stat().st_size > limits.max_output_bytes:
        raise ExportError("OUTPUT_LIMIT")
    return count


def _cleanup_xlsx_spools(workbook: Any) -> None:
    """Close and remove openpyxl's write-only worksheet spool files on abort."""

    for sheet in getattr(workbook, "worksheets", ()):
        writer = getattr(sheet, "_writer", None)
        if writer is None:
            continue
        try:
            # The worksheet API closes its row generator before closing the
            # XML stream, avoiding an unbalanced lxml context on cancellation.
            sheet.close()
        except Exception:
            pass
        try:
            writer.cleanup()
        except Exception:
            pass


def _xlsx_visible_cell(sheet: Any, value: Any) -> Any:
    if isinstance(value, str):
        visible = value if _utf16_units(value) <= _XLSX_CELL_LIMIT else "[full value in hidden typed metadata]"
    elif value is None:
        visible = ""
    elif isinstance(value, bool):
        return value
    elif isinstance(value, int):
        visible = str(value)
    elif isinstance(value, Decimal):
        visible = str(value)
    elif isinstance(value, float):
        visible = value.hex()
    elif isinstance(value, bytes):
        encoded = base64.b64encode(value).decode("ascii")
        visible = "base64:" + encoded if len(encoded) <= _XLSX_CELL_LIMIT else "[full value in hidden typed metadata]"
    elif isinstance(value, (_datetime.date, _datetime.time, _datetime.datetime)):
        visible = value.isoformat()
    else:
        raise ExportError("UNSUPPORTED_VALUE")
    if not isinstance(visible, str):
        visible = str(visible)
    if _utf16_units(visible) > _XLSX_CELL_LIMIT:
        visible = "[full value in hidden typed metadata]"
    return _safe_excel_string(sheet, visible)


def _utf16_units(value: str) -> int:
    try:
        return len(value.encode("utf-16-le")) // 2
    except UnicodeError as error:
        raise ExportError("INVALID_UNICODE") from error


def _split_xlsx_chunks(value: str) -> tuple[str, ...]:
    chunks: list[str] = []
    start = 0
    units = 0
    for index, character in enumerate(value):
        character_units = 2 if ord(character) > 0xFFFF else 1
        if units + character_units > _XLSX_META_CHUNK_CHARS:
            chunks.append(value[start:index])
            start = index
            units = 0
        units += character_units
    chunks.append(value[start:])
    return tuple(chunks)


def _safe_excel_string(sheet: Any, value: str) -> Any:
    try:
        from openpyxl.cell import WriteOnlyCell
    except Exception as error:
        raise ExportError("XLSX_UNAVAILABLE") from error
    cell = WriteOnlyCell(sheet, value=value)
    cell.data_type = "s"
    return cell


def _json_text(value: Any) -> str:
    try:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    except (TypeError, ValueError) as error:
        raise ExportError("UNSUPPORTED_VALUE") from error


def _json_bytes(value: Any) -> bytes:
    return _json_text(value).encode("utf-8")


class _Utf8BudgetWriter:
    def __init__(self, binary, max_bytes: int):
        self._binary = binary
        self._max_bytes = max_bytes
        self.bytes_written = 0

    def write(self, value: str) -> int:
        data = value.encode("utf-8")
        if self.bytes_written + len(data) > self._max_bytes:
            raise ExportError("OUTPUT_LIMIT")
        self._binary.write(data)
        self.bytes_written += len(data)
        return len(value)

    def flush(self) -> None:
        self._binary.flush()


def _fsync_file(path: Path) -> None:
    try:
        with path.open("r+b") as handle:
            handle.flush()
            os.fsync(handle.fileno())
    except OSError as error:
        raise ExportError("FSYNC_FAILED") from error


def _fsync_directory(path: Path) -> None:
    if os.name == "nt":
        return
    descriptor: int | None = None
    try:
        descriptor = os.open(path, os.O_RDONLY)
        os.fsync(descriptor)
    except OSError:
        # The complete file is already atomically published. Some filesystems
        # do not support opening or syncing directories, so this is best-effort.
        return
    finally:
        if descriptor is not None:
            os.close(descriptor)
