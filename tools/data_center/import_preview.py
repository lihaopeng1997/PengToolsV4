"""Bounded CSV/XLSX inspection and typed import planning.

This module deliberately stops before database work. It reads files only when
an explicit inspection or chunk-iteration function is called; constructing
sources and plans has no filesystem, database, UI, or configuration effects.
"""

from __future__ import annotations

import codecs
import csv
import datetime as datetime_module
import hashlib
import io
import os
import posixpath
import re
import stat
import tempfile
import threading
import zipfile
import xml.etree.ElementTree as element_tree
from dataclasses import dataclass, field, replace
from decimal import Decimal, InvalidOperation, localcontext
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Iterator, Sequence


class ImportErrorCode(str, Enum):
    INVALID_OPTIONS = "INVALID_OPTIONS"
    INVALID_SOURCE = "INVALID_SOURCE"
    INVALID_FORMAT = "INVALID_FORMAT"
    INVALID_ENCODING = "INVALID_ENCODING"
    INVALID_SHEET = "INVALID_SHEET"
    INVALID_MAPPING = "INVALID_MAPPING"
    MALFORMED_ROW = "MALFORMED_ROW"
    FORMULA_NOT_ALLOWED = "FORMULA_NOT_ALLOWED"
    CONVERSION_ERROR = "CONVERSION_ERROR"
    NULL_NOT_ALLOWED = "NULL_NOT_ALLOWED"
    LIMIT_EXCEEDED = "LIMIT_EXCEEDED"
    CANCELLED = "CANCELLED"
    CANCEL_CHECK_FAILED = "CANCEL_CHECK_FAILED"
    SOURCE_CHANGED = "SOURCE_CHANGED"
    APPROVAL_REQUIRED = "APPROVAL_REQUIRED"
    PREVIEW_HAS_ERRORS = "PREVIEW_HAS_ERRORS"


class ImportPreviewError(Exception):
    """Safe machine-readable failure; never contains a path or cell value."""

    def __init__(
        self,
        code: ImportErrorCode | str,
        *,
        row_index: int | None = None,
        column_index: int | None = None,
    ) -> None:
        try:
            selected_code = ImportErrorCode(code)
        except (ValueError, TypeError):
            selected_code = ImportErrorCode.INVALID_FORMAT
        self.code = selected_code.value
        self.row_index = row_index
        self.column_index = column_index
        super().__init__(self.code)


class ImportValueKind(str, Enum):
    TEXT = "text"
    INTEGER = "integer"
    DECIMAL = "decimal"
    DATE = "date"
    DATETIME = "datetime"


class RawCellKind(str, Enum):
    TEXT = "text"
    BLANK = "blank"
    NUMBER = "number"
    DATE = "date"
    DATETIME = "datetime"
    BOOLEAN = "boolean"
    UNSUPPORTED = "unsupported"


@dataclass(frozen=True, slots=True)
class CsvSource:
    path: Path | str | os.PathLike[str]
    encoding: str
    delimiter: str
    header_row_index: int = 0
    quotechar: str = '"'

    def __post_init__(self) -> None:
        try:
            normalized_path = Path(self.path)
            codecs.lookup(self.encoding)
        except (TypeError, ValueError, LookupError):
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS) from None
        if (
            "\x00" in str(normalized_path)
            or not isinstance(self.encoding, str)
            or not self.encoding
            or not isinstance(self.delimiter, str)
            or len(self.delimiter) != 1
            or self.delimiter in "\r\n"
            or not isinstance(self.quotechar, str)
            or len(self.quotechar) != 1
            or self.quotechar in "\r\n"
            or isinstance(self.header_row_index, bool)
            or not isinstance(self.header_row_index, int)
            or self.header_row_index < 0
        ):
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
        object.__setattr__(self, "path", normalized_path)


@dataclass(frozen=True, slots=True)
class XlsxSource:
    path: Path | str | os.PathLike[str]
    sheet_name: str
    header_row_index: int = 0

    def __post_init__(self) -> None:
        try:
            normalized_path = Path(self.path)
        except (TypeError, ValueError):
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS) from None
        if (
            "\x00" in str(normalized_path)
            or not isinstance(self.sheet_name, str)
            or not self.sheet_name
            or isinstance(self.header_row_index, bool)
            or not isinstance(self.header_row_index, int)
            or self.header_row_index < 0
        ):
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
        object.__setattr__(self, "path", normalized_path)


ImportSource = CsvSource | XlsxSource
CancelCheck = Callable[[], bool]


@dataclass(frozen=True, slots=True)
class ImportLimits:
    max_input_bytes: int = 32 * 1024 * 1024
    max_xlsx_entries: int = 2_048
    max_xlsx_uncompressed_bytes: int = 128 * 1024 * 1024
    max_xlsx_entry_bytes: int = 32 * 1024 * 1024
    max_xlsx_compression_ratio: int = 200
    max_columns: int = 512
    max_rows: int = 250_000
    max_cells: int = 2_000_000
    max_preview_rows: int = 100
    max_chunk_rows: int = 1_000
    max_cell_characters: int = 131_072

    def __post_init__(self) -> None:
        values = (
            self.max_input_bytes,
            self.max_xlsx_entries,
            self.max_xlsx_uncompressed_bytes,
            self.max_xlsx_entry_bytes,
            self.max_xlsx_compression_ratio,
            self.max_columns,
            self.max_rows,
            self.max_cells,
            self.max_preview_rows,
            self.max_chunk_rows,
            self.max_cell_characters,
        )
        if any(isinstance(value, bool) or not isinstance(value, int) or value <= 0 for value in values):
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
        hard_limits = (
            64 * 1024 * 1024,
            4_096,
            256 * 1024 * 1024,
            64 * 1024 * 1024,
            1_000,
            16_384,
            1_048_576,
            8_000_000,
            1_000,
            5_000,
            1_000_000,
        )
        if any(value > maximum for value, maximum in zip(values, hard_limits)):
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)


DEFAULT_IMPORT_LIMITS = ImportLimits()


@dataclass(frozen=True, slots=True)
class TargetColumn:
    name: str
    kind: ImportValueKind
    nullable: bool
    decimal_precision: int | None = None
    decimal_scale: int | None = None

    def __post_init__(self) -> None:
        try:
            selected_kind = ImportValueKind(self.kind)
        except (TypeError, ValueError):
            raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING) from None
        if not isinstance(self.name, str) or not self.name or "\x00" in self.name:
            raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
        if not isinstance(self.nullable, bool):
            raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
        has_decimal_meta = self.decimal_precision is not None or self.decimal_scale is not None
        if selected_kind is ImportValueKind.DECIMAL:
            if (
                isinstance(self.decimal_precision, bool)
                or not isinstance(self.decimal_precision, int)
                or isinstance(self.decimal_scale, bool)
                or not isinstance(self.decimal_scale, int)
                or self.decimal_precision <= 0
                or self.decimal_scale < 0
                or self.decimal_scale > self.decimal_precision
                or self.decimal_precision > 1_000
            ):
                raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
        elif has_decimal_meta:
            raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
        object.__setattr__(self, "kind", selected_kind)


@dataclass(frozen=True, slots=True)
class ColumnMapping:
    source_index: int
    target_name: str
    date_format: str | None = None

    def __post_init__(self) -> None:
        if (
            isinstance(self.source_index, bool)
            or not isinstance(self.source_index, int)
            or self.source_index < 0
            or not isinstance(self.target_name, str)
            or not self.target_name
            or "\x00" in self.target_name
            or (self.date_format is not None and (not isinstance(self.date_format, str) or not self.date_format))
        ):
            raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)


@dataclass(frozen=True, slots=True)
class NullRule:
    csv_null_tokens: tuple[str, ...]
    csv_empty_is_null: bool
    xlsx_null_tokens: tuple[str, ...]
    xlsx_empty_is_null: bool
    xlsx_blank_is_null: bool

    def __post_init__(self) -> None:
        try:
            csv_tokens = tuple(self.csv_null_tokens)
            xlsx_tokens = tuple(self.xlsx_null_tokens)
        except TypeError:
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS) from None
        if (
            any(not isinstance(token, str) for token in (*csv_tokens, *xlsx_tokens))
            or isinstance(self.csv_null_tokens, str)
            or isinstance(self.xlsx_null_tokens, str)
            or not isinstance(self.csv_empty_is_null, bool)
            or not isinstance(self.xlsx_empty_is_null, bool)
            or not isinstance(self.xlsx_blank_is_null, bool)
        ):
            raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
        object.__setattr__(self, "csv_null_tokens", csv_tokens)
        object.__setattr__(self, "xlsx_null_tokens", xlsx_tokens)


@dataclass(frozen=True, slots=True)
class SourceFingerprint:
    size_bytes: int
    modified_ns: int
    device: int
    inode: int
    sha256: str


class _SourceSnapshot:
    """Private bounded copy used as the only parser input for an approval."""

    __slots__ = ("path", "_closed", "_close_requested", "_started", "_active", "_lock")

    def __init__(self, path: Path) -> None:
        self.path = path
        self._closed = False
        self._close_requested = False
        self._started = False
        self._active = 0
        self._lock = threading.Lock()

    @property
    def closed(self) -> bool:
        with self._lock:
            return self._closed or self._close_requested

    def close(self) -> None:
        with self._lock:
            self._close_requested = True
            if self._closed or self._active:
                return
            self._closed = True
        self._delete_owned_file()

    def begin_import(self) -> Path:
        with self._lock:
            if self._closed or self._close_requested or self._started:
                raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)
            self._started = True
            self._active += 1
            return self.path

    def end_import(self) -> None:
        with self._lock:
            self._close_requested = True
            if self._active:
                self._active -= 1
            should_delete = self._active == 0 and not self._closed
            if should_delete:
                self._closed = True
        if should_delete:
            self._delete_owned_file()

    def _delete_owned_file(self) -> None:
        try:
            os.chmod(self.path, stat.S_IREAD | stat.S_IWRITE)
        except OSError:
            pass
        try:
            self.path.unlink(missing_ok=True)
        except OSError:
            pass

    def __del__(self) -> None:
        try:
            self.close()
        except Exception:
            pass


@dataclass(frozen=True, slots=True)
class SourceColumn:
    index: int
    label: str


@dataclass(frozen=True, slots=True)
class RawCell:
    value: Any
    kind: RawCellKind


@dataclass(frozen=True, slots=True)
class SourceRow:
    row_index: int
    cells: tuple[RawCell, ...]


@dataclass(frozen=True, slots=True)
class SourcePreview:
    source: ImportSource
    fingerprint: SourceFingerprint
    columns: tuple[SourceColumn, ...]
    rows: tuple[SourceRow, ...]
    has_more: bool
    limits: ImportLimits
    _snapshot: _SourceSnapshot | None = field(default=None, repr=False, compare=False)


@dataclass(frozen=True, slots=True)
class ImportIssue:
    code: str
    row_index: int
    column_index: int


@dataclass(frozen=True, slots=True)
class ImportPreviewRow:
    source_row_index: int
    values: tuple[Any, ...]


@dataclass(frozen=True, slots=True)
class ImportPreview:
    source_preview: SourcePreview
    targets: tuple[TargetColumn, ...]
    mapping: tuple[ColumnMapping, ...]
    null_rule: NullRule
    rows: tuple[ImportPreviewRow, ...]
    issues: tuple[ImportIssue, ...]


@dataclass(frozen=True, slots=True)
class ApprovedImport:
    source: ImportSource
    fingerprint: SourceFingerprint
    columns: tuple[SourceColumn, ...]
    targets: tuple[TargetColumn, ...]
    mapping: tuple[ColumnMapping, ...]
    null_rule: NullRule
    limits: ImportLimits
    _snapshot: _SourceSnapshot | None = field(default=None, repr=False, compare=False)


@dataclass(frozen=True, slots=True)
class ImportChunk:
    start_row_index: int
    target_names: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]


@dataclass(frozen=True, slots=True)
class _XlsxMetadata:
    sheet_entry: str
    max_row: int
    header_width: int
    data_row_count: int


@dataclass(frozen=True, slots=True)
class _XlsxRowInfo:
    row_number: int
    numeric_values: dict[int, str]
    formula_columns: frozenset[int]


_CSV_LIMIT_LOCK = threading.RLock()
_DECIMAL_TEXT_RE = re.compile(r"^[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?$")
_INTEGER_TEXT_RE = re.compile(r"^[+-]?[0-9]+$")
_CELL_REFERENCE_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$")
_XML_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_XML_DOC_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_XML_PACKAGE_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def inspect_import_source(
    source: ImportSource,
    *,
    limits: ImportLimits = DEFAULT_IMPORT_LIMITS,
    cancel_check: CancelCheck | None = None,
) -> SourcePreview:
    """Read a bounded header/sample and hash the complete bounded source file."""

    if not isinstance(limits, ImportLimits):
        raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
    if not isinstance(source, (CsvSource, XlsxSource)):
        raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
    if source.header_row_index >= limits.max_rows:
        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
    fingerprint, snapshot = _capture_source(
        source.path,
        ".csv" if isinstance(source, CsvSource) else ".xlsx",
        limits,
        cancel_check,
    )
    try:
        snapshot_source = replace(source, path=snapshot.path)
        if isinstance(snapshot_source, CsvSource):
            preview = _inspect_csv(snapshot_source, fingerprint, limits, cancel_check)
        else:
            metadata = _validate_xlsx_archive(snapshot_source, limits, cancel_check)
            _verify_fingerprint(source.path, fingerprint, limits, cancel_check)
            _verify_snapshot(snapshot.path, fingerprint, limits, cancel_check)
            preview = _inspect_xlsx(snapshot_source, fingerprint, limits, metadata, cancel_check)
        _verify_snapshot(snapshot.path, fingerprint, limits, cancel_check)
        _verify_fingerprint(source.path, fingerprint, limits, cancel_check)
        return replace(preview, source=source, _snapshot=snapshot)
    except Exception:
        snapshot.close()
        raise


def preview_import(
    source_preview: SourcePreview,
    mapping: Sequence[ColumnMapping],
    targets: Sequence[TargetColumn],
    null_rule: NullRule,
) -> ImportPreview:
    """Apply explicit mappings and conversion rules to the finite source sample."""

    if not isinstance(source_preview, SourcePreview) or not isinstance(null_rule, NullRule):
        raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
    normalized_mapping, normalized_targets = _validate_mapping(source_preview.columns, mapping, targets)
    result_rows: list[ImportPreviewRow] = []
    issues: list[ImportIssue] = []
    for row in source_preview.rows:
        converted: list[Any] = []
        for column_mapping, target in zip(normalized_mapping, normalized_targets):
            raw_cell = row.cells[column_mapping.source_index]
            try:
                value = _convert_cell(
                    raw_cell,
                    target,
                    column_mapping,
                    null_rule,
                    is_xlsx=isinstance(source_preview.source, XlsxSource),
                    row_index=row.row_index,
                    column_index=column_mapping.source_index,
                )
            except ImportPreviewError as error:
                issues.append(ImportIssue(error.code, row.row_index, column_mapping.source_index))
                converted.append(None)
            else:
                converted.append(value)
        result_rows.append(ImportPreviewRow(row.row_index, tuple(converted)))
    return ImportPreview(
        source_preview=source_preview,
        targets=normalized_targets,
        mapping=normalized_mapping,
        null_rule=null_rule,
        rows=tuple(result_rows),
        issues=tuple(issues),
    )


def approve_import(preview: ImportPreview, *, user_approved: bool) -> ApprovedImport:
    """Create an immutable file/mapping plan after a separate explicit approval."""

    if not isinstance(preview, ImportPreview):
        raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
    if user_approved is not True:
        raise ImportPreviewError(ImportErrorCode.APPROVAL_REQUIRED)
    if preview.issues:
        raise ImportPreviewError(ImportErrorCode.PREVIEW_HAS_ERRORS)
    source_preview = preview.source_preview
    if source_preview._snapshot is None or source_preview._snapshot.closed:
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)
    return ApprovedImport(
        source=source_preview.source,
        fingerprint=source_preview.fingerprint,
        columns=source_preview.columns,
        targets=preview.targets,
        mapping=preview.mapping,
        null_rule=preview.null_rule,
        limits=source_preview.limits,
        _snapshot=source_preview._snapshot,
    )


def close_import_snapshot(value: SourcePreview | ImportPreview | ApprovedImport) -> None:
    """Request deletion of the private copy after any active reader releases it."""

    if isinstance(value, SourcePreview):
        snapshot = value._snapshot
    elif isinstance(value, ImportPreview):
        snapshot = value.source_preview._snapshot
    elif isinstance(value, ApprovedImport):
        snapshot = value._snapshot
    else:
        raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
    if snapshot is not None:
        snapshot.close()


def iter_import_chunks(
    approved: ApprovedImport,
    *,
    chunk_rows: int | None = None,
    cancel_check: CancelCheck | None = None,
) -> Iterator[ImportChunk]:
    """Lazily convert rows in bounded chunks; performs no database writes.

    Parsing uses the private byte-for-byte snapshot that was hashed during
    inspection; the original path is rechecked before parsing, before yields,
    and after parsing. The snapshot supports one iterator consumption and is
    removed when that iterator ends or is closed. A mutation detected after
    earlier chunks were yielded requires the caller's dedicated import
    transaction to roll back; this iterator cannot undo consumer work.
    The caller must exhaust this iterator successfully before offering commit,
    roll back only that import transaction on any error/cancel, and commit only
    after the user's separate commit decision.
    """

    if not isinstance(approved, ApprovedImport):
        raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
    selected_chunk_rows = approved.limits.max_chunk_rows if chunk_rows is None else chunk_rows
    if (
        isinstance(selected_chunk_rows, bool)
        or not isinstance(selected_chunk_rows, int)
        or selected_chunk_rows <= 0
        or selected_chunk_rows > approved.limits.max_chunk_rows
    ):
        raise ImportPreviewError(ImportErrorCode.INVALID_OPTIONS)
    snapshot = approved._snapshot
    if snapshot is None:
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)
    snapshot_path = snapshot.begin_import()
    rows: Iterator[SourceRow] | None = None
    try:
        _verify_fingerprint(approved.source.path, approved.fingerprint, approved.limits, cancel_check)
        _verify_snapshot(snapshot_path, approved.fingerprint, approved.limits, cancel_check)
        _validate_mapping(approved.columns, approved.mapping, approved.targets)
        snapshot_source = replace(approved.source, path=snapshot_path)
        if isinstance(snapshot_source, CsvSource):
            rows = _iter_csv_data_rows(
                snapshot_source,
                approved.columns,
                approved.limits,
                cancel_check,
            )
        else:
            metadata = _validate_xlsx_archive(snapshot_source, approved.limits, cancel_check)
            _verify_snapshot(snapshot_path, approved.fingerprint, approved.limits, cancel_check)
            rows = _iter_xlsx_data_rows(
                snapshot_source,
                approved.columns,
                approved.limits,
                metadata,
                cancel_check,
            )
        mapping_targets = tuple(
            (column_mapping, next(target for target in approved.targets if target.name == column_mapping.target_name))
            for column_mapping in approved.mapping
        )
        pending: list[tuple[Any, ...]] = []
        pending_start: int | None = None
        for row in rows:
            _check_cancel(cancel_check)
            _check_path_identity(approved.source.path, approved.fingerprint)
            if row.row_index >= approved.limits.max_rows:
                raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED, row_index=row.row_index)
            converted: list[Any] = []
            for column_mapping, target in mapping_targets:
                try:
                    value = _convert_cell(
                        row.cells[column_mapping.source_index],
                        target,
                        column_mapping,
                        approved.null_rule,
                        is_xlsx=isinstance(approved.source, XlsxSource),
                        row_index=row.row_index,
                        column_index=column_mapping.source_index,
                    )
                except ImportPreviewError as error:
                    raise ImportPreviewError(
                        error.code,
                        row_index=error.row_index,
                        column_index=error.column_index,
                    ) from None
                converted.append(value)
            if pending_start is None:
                pending_start = row.row_index
            pending.append(tuple(converted))
            if len(pending) == selected_chunk_rows:
                _check_path_identity(approved.source.path, approved.fingerprint)
                yield ImportChunk(
                    start_row_index=pending_start,
                    target_names=tuple(target.name for _, target in mapping_targets),
                    rows=tuple(pending),
                )
                pending = []
                pending_start = None
        if pending:
            _check_path_identity(approved.source.path, approved.fingerprint)
            yield ImportChunk(
                start_row_index=pending_start if pending_start is not None else 0,
                target_names=tuple(target.name for _, target in mapping_targets),
                rows=tuple(pending),
            )
        # A partial final chunk is not successful exhaustion until the caller
        # resumes the iterator and these terminal checks pass.
        _check_cancel(cancel_check)
        _verify_snapshot(snapshot_path, approved.fingerprint, approved.limits, cancel_check)
        _verify_fingerprint(approved.source.path, approved.fingerprint, approved.limits, cancel_check)
    except ImportPreviewError:
        try:
            _verify_fingerprint(approved.source.path, approved.fingerprint, approved.limits, None)
        except ImportPreviewError as changed_error:
            if changed_error.code == ImportErrorCode.SOURCE_CHANGED.value:
                raise changed_error
        raise
    finally:
        try:
            if rows is not None:
                rows.close()
        finally:
            snapshot.end_import()


def _validate_mapping(
    columns: Sequence[SourceColumn],
    mapping: Sequence[ColumnMapping],
    targets: Sequence[TargetColumn],
) -> tuple[tuple[ColumnMapping, ...], tuple[TargetColumn, ...]]:
    try:
        source_columns = tuple(columns)
        selected_mapping = tuple(mapping)
        selected_targets = tuple(targets)
    except TypeError:
        raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING) from None
    if not selected_mapping or not selected_targets:
        raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
    if any(not isinstance(column, SourceColumn) for column in source_columns):
        raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
    if any(not isinstance(item, ColumnMapping) for item in selected_mapping):
        raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
    if any(not isinstance(item, TargetColumn) for item in selected_targets):
        raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
    target_names = tuple(target.name for target in selected_targets)
    mapping_targets = tuple(item.target_name for item in selected_mapping)
    source_indices = tuple(item.source_index for item in selected_mapping)
    if (
        len(set(target_names)) != len(target_names)
        or len(set(mapping_targets)) != len(mapping_targets)
        or len(set(source_indices)) != len(source_indices)
        or set(mapping_targets) != set(target_names)
        or any(index >= len(source_columns) for index in source_indices)
        or any(column.index != index for index, column in enumerate(source_columns))
    ):
        raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
    target_by_name = {target.name: target for target in selected_targets}
    for item in selected_mapping:
        target = target_by_name[item.target_name]
        if item.date_format is not None and target.kind not in (ImportValueKind.DATE, ImportValueKind.DATETIME):
            raise ImportPreviewError(ImportErrorCode.INVALID_MAPPING)
    targets_in_mapping_order = tuple(target_by_name[item.target_name] for item in selected_mapping)
    return selected_mapping, targets_in_mapping_order


def _convert_cell(
    raw_cell: RawCell,
    target: TargetColumn,
    mapping: ColumnMapping,
    null_rule: NullRule,
    *,
    is_xlsx: bool,
    row_index: int,
    column_index: int,
) -> Any:
    value = raw_cell.value
    is_null = False
    if raw_cell.kind is RawCellKind.BLANK:
        is_null = null_rule.xlsx_blank_is_null
        if not is_null and target.kind is ImportValueKind.TEXT:
            value = ""
    elif raw_cell.kind is RawCellKind.TEXT:
        tokens = null_rule.xlsx_null_tokens if is_xlsx else null_rule.csv_null_tokens
        is_null = value in tokens
        if value == "":
            is_null = is_null or (null_rule.xlsx_empty_is_null if is_xlsx else null_rule.csv_empty_is_null)
    if is_null:
        if not target.nullable:
            raise ImportPreviewError(
                ImportErrorCode.NULL_NOT_ALLOWED,
                row_index=row_index,
                column_index=column_index,
            )
        return None

    try:
        if target.kind is ImportValueKind.TEXT:
            if raw_cell.kind is RawCellKind.TEXT:
                return value
            if raw_cell.kind is RawCellKind.NUMBER:
                return value
            if raw_cell.kind is RawCellKind.BLANK and not null_rule.xlsx_blank_is_null:
                return ""
            raise ValueError

        if target.kind is ImportValueKind.INTEGER:
            if raw_cell.kind is RawCellKind.NUMBER:
                number = _parse_decimal(str(value))
                if number != number.to_integral_value():
                    raise ValueError
                if number and number.copy_abs().adjusted() >= 1_000:
                    raise ValueError
                return int(number)
            if raw_cell.kind is RawCellKind.TEXT and _INTEGER_TEXT_RE.fullmatch(value):
                return int(value)
            raise ValueError

        if target.kind is ImportValueKind.DECIMAL:
            if raw_cell.kind not in (RawCellKind.NUMBER, RawCellKind.TEXT):
                raise ValueError
            text_value = str(value)
            number = _parse_decimal(text_value)
            return _validate_decimal_precision(number, target)

        if target.kind is ImportValueKind.DATE:
            if raw_cell.kind is RawCellKind.DATE and isinstance(value, datetime_module.date):
                return value
            if raw_cell.kind is RawCellKind.DATETIME and isinstance(value, datetime_module.datetime):
                if value.time() != datetime_module.time(0, 0):
                    raise ValueError
                return value.date()
            if raw_cell.kind is RawCellKind.TEXT and mapping.date_format is not None:
                parsed = datetime_module.datetime.strptime(value, mapping.date_format)
                if parsed.strftime(mapping.date_format) != value:
                    raise ValueError
                if parsed.time() != datetime_module.time(0, 0):
                    raise ValueError
                return parsed.date()
            raise ValueError

        if target.kind is ImportValueKind.DATETIME:
            if raw_cell.kind is RawCellKind.DATETIME and isinstance(value, datetime_module.datetime):
                return value
            if raw_cell.kind is RawCellKind.TEXT and mapping.date_format is not None:
                parsed = datetime_module.datetime.strptime(value, mapping.date_format)
                if parsed.strftime(mapping.date_format) != value:
                    raise ValueError
                return parsed
            raise ValueError
    except ImportPreviewError:
        raise
    except (ValueError, TypeError, OverflowError, InvalidOperation):
        raise ImportPreviewError(
            ImportErrorCode.CONVERSION_ERROR,
            row_index=row_index,
            column_index=column_index,
        ) from None
    raise ImportPreviewError(
        ImportErrorCode.CONVERSION_ERROR,
        row_index=row_index,
        column_index=column_index,
    )


def _parse_decimal(text: str) -> Decimal:
    if not _DECIMAL_TEXT_RE.fullmatch(text) or len(text) > 1_000:
        raise ValueError
    value = Decimal(text)
    if (
        not value.is_finite()
        or (value and abs(value.adjusted()) > 1_000)
        or abs(value.as_tuple().exponent) > 1_000
    ):
        raise ValueError
    return value


def _validate_decimal_precision(value: Decimal, target: TargetColumn) -> Decimal:
    if target.decimal_precision is None or target.decimal_scale is None:
        raise ValueError
    with localcontext() as context:
        context.prec = max(40, len(value.as_tuple().digits) + target.decimal_precision + 4)
        quantum = Decimal(1).scaleb(-target.decimal_scale)
        normalized = value.quantize(quantum)
    if normalized != value:
        raise ValueError
    integer_digits = max(normalized.copy_abs().adjusted() + 1, 0) if normalized else 0
    if integer_digits > target.decimal_precision - target.decimal_scale:
        raise ValueError
    return normalized


def _inspect_csv(
    source: CsvSource,
    fingerprint: SourceFingerprint,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> SourcePreview:
    rows: list[SourceRow] = []
    has_more = False
    try:
        with _csv_limits(limits.max_cell_characters):
            with source.path.open("rb") as binary_stream:
                text_stream = io.TextIOWrapper(binary_stream, encoding=source.encoding, errors="strict", newline="")
                try:
                    reader = csv.reader(
                        text_stream,
                        delimiter=source.delimiter,
                        quotechar=source.quotechar,
                        strict=True,
                    )
                    header: list[str] | None = None
                    for index in range(source.header_row_index + 1):
                        _check_cancel(cancel_check)
                        header = next(reader, None)
                        if header is None:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                    if not header or len(header) > limits.max_columns:
                        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                    _validate_cell_lengths(header, limits)
                    columns = tuple(SourceColumn(index, label) for index, label in enumerate(header))
                    cells_seen = len(header)
                    for sample_index in range(limits.max_preview_rows + 1):
                        _check_cancel(cancel_check)
                        values = next(reader, None)
                        if values is None:
                            break
                        row_index = source.header_row_index + sample_index + 1
                        if row_index >= limits.max_rows:
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED, row_index=row_index)
                        if len(values) != len(header):
                            raise ImportPreviewError(ImportErrorCode.MALFORMED_ROW, row_index=row_index)
                        _validate_cell_lengths(values, limits, row_index=row_index)
                        cells_seen += len(values)
                        if cells_seen > limits.max_cells:
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED, row_index=row_index)
                        if sample_index == limits.max_preview_rows:
                            has_more = True
                        else:
                            rows.append(
                                SourceRow(
                                    row_index=row_index,
                                    cells=tuple(RawCell(value, RawCellKind.TEXT) for value in values),
                                )
                            )
                finally:
                    try:
                        text_stream.detach()
                    except (ValueError, OSError):
                        pass
    except ImportPreviewError:
        raise
    except UnicodeError:
        raise ImportPreviewError(ImportErrorCode.INVALID_ENCODING) from None
    except (csv.Error, OSError, ValueError):
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None
    if header is None:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
    return SourcePreview(
        source=source,
        fingerprint=fingerprint,
        columns=columns,
        rows=tuple(rows),
        has_more=has_more,
        limits=limits,
    )


def _csv_limits(max_field_size: int):
    class _CsvLimitContext:
        def __enter__(self):
            _CSV_LIMIT_LOCK.acquire()
            self.previous = csv.field_size_limit()
            csv.field_size_limit(max_field_size)

        def __exit__(self, exc_type, exc_value, traceback):
            csv.field_size_limit(self.previous)
            _CSV_LIMIT_LOCK.release()

    return _CsvLimitContext()


def _inspect_xlsx(
    source: XlsxSource,
    fingerprint: SourceFingerprint,
    limits: ImportLimits,
    metadata: _XlsxMetadata,
    cancel_check: CancelCheck | None,
) -> SourcePreview:
    worksheet, workbook = _open_xlsx_worksheet(source, metadata)
    rows: list[SourceRow] = []
    columns: tuple[SourceColumn, ...] | None = None
    has_more = False
    try:
        raw_rows = _iter_xlsx_raw_rows(source.path, metadata.sheet_entry, limits, cancel_check)
        raw_info = next(raw_rows, None)
        for row_number, cells in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=metadata.max_row,
                min_col=1,
                max_col=metadata.header_width,
                values_only=False,
            ),
            start=1,
        ):
            _check_cancel(cancel_check)
            if (
                row_number > source.header_row_index + 1
                and len(rows) >= limits.max_preview_rows
            ):
                has_more = True
                break
            info = _take_xlsx_row_info(raw_rows, raw_info, row_number)
            raw_info = info[1]
            row_info = info[0]
            raw_cells = _xlsx_cells(cells, row_info, row_number, limits)
            raw_cells = _pad_xlsx_row(raw_cells, metadata.header_width, row_number)
            if row_number == source.header_row_index + 1:
                labels: list[str] = []
                for cell in raw_cells:
                    if cell.kind is RawCellKind.TEXT:
                        labels.append(cell.value)
                    elif cell.kind is RawCellKind.BLANK:
                        labels.append("")
                    elif cell.kind is RawCellKind.NUMBER:
                        labels.append(cell.value)
                    else:
                        raise ImportPreviewError(
                            ImportErrorCode.INVALID_FORMAT,
                            row_index=row_number - 1,
                        )
                columns = tuple(SourceColumn(index, label) for index, label in enumerate(labels))
                continue
            if row_number <= source.header_row_index + 1:
                continue
            rows.append(SourceRow(row_number - 1, raw_cells))
    except ImportPreviewError:
        raise
    except (OSError, zipfile.BadZipFile, KeyError, ValueError, TypeError):
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None
    finally:
        if "raw_rows" in locals():
            raw_rows.close()
        workbook.close()
    if columns is None:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
    return SourcePreview(
        source=source,
        fingerprint=fingerprint,
        columns=columns,
        rows=tuple(rows),
        has_more=has_more or metadata.data_row_count > len(rows),
        limits=limits,
    )


def _iter_csv_data_rows(
    source: CsvSource,
    columns: Sequence[SourceColumn],
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> Iterator[SourceRow]:
    seen_rows = 0
    cells_seen = len(columns)
    try:
        with _csv_limits(limits.max_cell_characters):
            with source.path.open("rb") as binary_stream:
                text_stream = io.TextIOWrapper(binary_stream, encoding=source.encoding, errors="strict", newline="")
                try:
                    reader = csv.reader(
                        text_stream,
                        delimiter=source.delimiter,
                        quotechar=source.quotechar,
                        strict=True,
                    )
                    for _ in range(source.header_row_index + 1):
                        _check_cancel(cancel_check)
                        if next(reader, None) is None:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                    for values in reader:
                        _check_cancel(cancel_check)
                        row_index = source.header_row_index + seen_rows + 1
                        seen_rows += 1
                        if row_index >= limits.max_rows:
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED, row_index=row_index)
                        if len(values) != len(columns):
                            raise ImportPreviewError(ImportErrorCode.MALFORMED_ROW, row_index=row_index)
                        _validate_cell_lengths(values, limits, row_index=row_index)
                        cells_seen += len(values)
                        if cells_seen > limits.max_cells:
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED, row_index=row_index)
                        yield SourceRow(
                            row_index,
                            tuple(RawCell(value, RawCellKind.TEXT) for value in values),
                        )
                finally:
                    try:
                        text_stream.detach()
                    except (ValueError, OSError):
                        pass
    except ImportPreviewError:
        raise
    except UnicodeError:
        raise ImportPreviewError(ImportErrorCode.INVALID_ENCODING) from None
    except (csv.Error, OSError, ValueError):
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None


def _iter_xlsx_data_rows(
    source: XlsxSource,
    columns: Sequence[SourceColumn],
    limits: ImportLimits,
    metadata: _XlsxMetadata,
    cancel_check: CancelCheck | None,
) -> Iterator[SourceRow]:
    worksheet, workbook = _open_xlsx_worksheet(source, metadata)
    seen_rows = 0
    raw_rows = _iter_xlsx_raw_rows(source.path, metadata.sheet_entry, limits, cancel_check)
    try:
        raw_info = next(raw_rows, None)
        for row_number, cells in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=metadata.max_row,
                min_col=1,
                max_col=len(columns),
                values_only=False,
            ),
            start=1,
        ):
            _check_cancel(cancel_check)
            row_info, raw_info = _take_xlsx_row_info(raw_rows, raw_info, row_number)
            raw_cells = _xlsx_cells(cells, row_info, row_number, limits)
            raw_cells = _pad_xlsx_row(raw_cells, len(columns), row_number)
            if row_number <= source.header_row_index + 1:
                if row_number == source.header_row_index + 1:
                    if any(index in row_info.formula_columns for index in range(len(columns))):
                        raise ImportPreviewError(ImportErrorCode.FORMULA_NOT_ALLOWED, row_index=row_number - 1)
                continue
            seen_rows += 1
            if seen_rows > limits.max_rows:
                raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED, row_index=row_number - 1)
            if len(raw_cells) != len(columns):
                raise ImportPreviewError(ImportErrorCode.MALFORMED_ROW, row_index=row_number - 1)
            if row_info.formula_columns:
                formula_column = min(row_info.formula_columns)
                raise ImportPreviewError(
                    ImportErrorCode.FORMULA_NOT_ALLOWED,
                    row_index=row_number - 1,
                    column_index=formula_column,
                )
            yield SourceRow(row_number - 1, raw_cells)
    except ImportPreviewError:
        raise
    except (OSError, zipfile.BadZipFile, KeyError, ValueError, TypeError):
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None
    finally:
        raw_rows.close()
        workbook.close()


def _open_xlsx_worksheet(source: XlsxSource, metadata: _XlsxMetadata):
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(source.path, read_only=True, data_only=False, keep_links=False)
    except ImportPreviewError:
        raise
    except Exception:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None
    if source.sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ImportPreviewError(ImportErrorCode.INVALID_SHEET)
    return workbook[source.sheet_name], workbook


def _xlsx_cells(
    cells: Sequence[Any],
    info: _XlsxRowInfo,
    row_number: int,
    limits: ImportLimits,
) -> tuple[RawCell, ...]:
    result: list[RawCell] = []
    for column_index, cell in enumerate(cells):
        if cell.data_type == "f" or column_index in info.formula_columns:
            raise ImportPreviewError(
                ImportErrorCode.FORMULA_NOT_ALLOWED,
                row_index=row_number - 1,
                column_index=column_index,
            )
        if cell.data_type == "e":
            raise ImportPreviewError(
                ImportErrorCode.INVALID_FORMAT,
                row_index=row_number - 1,
                column_index=column_index,
            )
        value = cell.value
        if value is None:
            kind = RawCellKind.BLANK
        elif isinstance(value, bool):
            kind = RawCellKind.BOOLEAN
        elif isinstance(value, datetime_module.datetime):
            kind = RawCellKind.DATETIME
        elif isinstance(value, datetime_module.date):
            kind = RawCellKind.DATE
        elif isinstance(value, str):
            kind = RawCellKind.TEXT
        elif isinstance(value, (int, float, Decimal)):
            if cell.data_type == "n":
                lexical = info.numeric_values.get(column_index)
                if lexical is None:
                    raise ImportPreviewError(
                        ImportErrorCode.INVALID_FORMAT,
                        row_index=row_number - 1,
                        column_index=column_index,
                    )
                value = lexical
            kind = RawCellKind.NUMBER
        else:
            kind = RawCellKind.UNSUPPORTED
        if isinstance(value, str) and len(value) > limits.max_cell_characters:
            raise ImportPreviewError(
                ImportErrorCode.LIMIT_EXCEEDED,
                row_index=row_number - 1,
                column_index=column_index,
            )
        result.append(RawCell(value, kind))
    return tuple(result)


def _pad_xlsx_row(
    cells: tuple[RawCell, ...],
    header_width: int,
    row_number: int,
) -> tuple[RawCell, ...]:
    """Represent omitted trailing worksheet cells as explicit blank values."""

    if len(cells) > header_width:
        raise ImportPreviewError(ImportErrorCode.MALFORMED_ROW, row_index=row_number - 1)
    if len(cells) == header_width:
        return cells
    return cells + (RawCell(None, RawCellKind.BLANK),) * (header_width - len(cells))


def _take_xlsx_row_info(raw_rows, current: _XlsxRowInfo | None, row_number: int):
    info = current
    while info is not None and info.row_number < row_number:
        info = next(raw_rows, None)
    if info is None or info.row_number != row_number:
        return _XlsxRowInfo(row_number, {}, frozenset()), info
    return info, next(raw_rows, None)


def _iter_xlsx_raw_rows(
    path: Path,
    sheet_entry: str,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> Iterator[_XlsxRowInfo]:
    try:
        with zipfile.ZipFile(path, "r") as archive:
            with archive.open(sheet_entry, "r") as stream:
                current_row: int | None = None
                numeric_values: dict[int, str] = {}
                formula_columns: set[int] = set()
                for event, element in element_tree.iterparse(stream, events=("start", "end")):
                    _check_cancel(cancel_check)
                    local_name = _local_name(element.tag)
                    if event == "start" and local_name == "row":
                        current_row = _positive_xml_integer(element.attrib.get("r"))
                        numeric_values = {}
                        formula_columns = set()
                    elif event == "end" and local_name == "c":
                        cell_ref = _parse_cell_reference(element.attrib.get("r"))
                        if cell_ref is None:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                        column_index, cell_row = cell_ref
                        if current_row != cell_row:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                        if any(_local_name(child.tag) == "f" for child in element):
                            formula_columns.add(column_index)
                        cell_type = element.attrib.get("t")
                        if cell_type in (None, "n"):
                            value_element = next(
                                (child for child in element if _local_name(child.tag) == "v"),
                                None,
                            )
                            if value_element is not None and value_element.text is not None:
                                numeric_values[column_index] = value_element.text
                        element.clear()
                    elif event == "end" and local_name == "row":
                        if current_row is None:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                        yield _XlsxRowInfo(
                            current_row,
                            numeric_values,
                            frozenset(formula_columns),
                        )
                        element.clear()
                        current_row = None
    except ImportPreviewError:
        raise
    except (OSError, zipfile.BadZipFile, KeyError, element_tree.ParseError):
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None


def _validate_xlsx_archive(
    source: XlsxSource,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> _XlsxMetadata:
    try:
        with zipfile.ZipFile(source.path, "r") as archive:
            infos = archive.infolist()
            if not infos or len(infos) > limits.max_xlsx_entries:
                raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
            names: set[str] = set()
            total_declared = 0
            for info in infos:
                _check_cancel(cancel_check)
                if (
                    info.filename in names
                    or not _safe_zip_member(info.filename)
                    or info.flag_bits & 0x1
                ):
                    raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                names.add(info.filename)
                if info.file_size < 0 or info.compress_size < 0:
                    raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                total_declared += info.file_size
                if (
                    info.file_size > limits.max_xlsx_entry_bytes
                    or total_declared > limits.max_xlsx_uncompressed_bytes
                    or (
                        info.file_size > 0
                        and info.file_size / max(info.compress_size, 1) > limits.max_xlsx_compression_ratio
                    )
                ):
                    raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
            required = {
                "[Content_Types].xml",
                "xl/workbook.xml",
                "xl/_rels/workbook.xml.rels",
            }
            if not required.issubset(names):
                raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
            total_actual = 0
            for info in infos:
                with archive.open(info, "r") as entry:
                    actual_entry = 0
                    while True:
                        _check_cancel(cancel_check)
                        block = entry.read(64 * 1024)
                        if not block:
                            break
                        actual_entry += len(block)
                        total_actual += len(block)
                        if (
                            actual_entry > limits.max_xlsx_entry_bytes
                            or total_actual > limits.max_xlsx_uncompressed_bytes
                        ):
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                    if actual_entry != info.file_size:
                        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
            workbook_root = _read_xml_member(archive, "xl/workbook.xml", limits)
            relationships_root = _read_xml_member(archive, "xl/_rels/workbook.xml.rels", limits)
            relationship_targets = {}
            for relationship in relationships_root.findall(f"{_XML_PACKAGE_REL_NS}Relationship"):
                relation_id = relationship.attrib.get("Id")
                target = relationship.attrib.get("Target")
                if relation_id and target and relationship.attrib.get("TargetMode") != "External":
                    relationship_targets[relation_id] = _resolve_sheet_target(target)
            sheet_entry = None
            for sheet in workbook_root.findall(f"{_XML_MAIN_NS}sheets/{_XML_MAIN_NS}sheet"):
                if sheet.attrib.get("name") == source.sheet_name:
                    relation_id = sheet.attrib.get(f"{_XML_DOC_REL_NS}id")
                    sheet_entry = relationship_targets.get(relation_id)
                    break
            if sheet_entry is None or sheet_entry not in names:
                raise ImportPreviewError(ImportErrorCode.INVALID_SHEET)
            metadata = _scan_xlsx_sheet_metadata(
                archive,
                sheet_entry,
                source.header_row_index,
                limits,
                cancel_check,
            )
            return _XlsxMetadata(sheet_entry, metadata[0], metadata[1], metadata[2])
    except ImportPreviewError:
        raise
    except Exception:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None


def _read_xml_member(archive: zipfile.ZipFile, member: str, limits: ImportLimits):
    try:
        info = archive.getinfo(member)
        if info.file_size > limits.max_xlsx_entry_bytes:
            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
        with archive.open(info, "r") as stream:
            content = stream.read(limits.max_xlsx_entry_bytes + 1)
    except ImportPreviewError:
        raise
    except (KeyError, OSError, zipfile.BadZipFile, RuntimeError):
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None
    if len(content) > limits.max_xlsx_entry_bytes:
        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
    upper_content = content.upper()
    if b"<!DOCTYPE" in upper_content or b"<!ENTITY" in upper_content:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
    return element_tree.fromstring(content)


def _scan_xlsx_sheet_metadata(
    archive: zipfile.ZipFile,
    sheet_entry: str,
    header_row_index: int,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> tuple[int, int, int]:
    try:
        with archive.open(sheet_entry, "r") as stream:
            xml_content = stream.read(limits.max_xlsx_entry_bytes + 1)
    except (KeyError, OSError, zipfile.BadZipFile, RuntimeError):
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None
    if len(xml_content) > limits.max_xlsx_entry_bytes:
        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
    upper_content = xml_content.upper()
    if b"<!DOCTYPE" in upper_content or b"<!ENTITY" in upper_content:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)

    declared_max_row = 0
    declared_max_column = 0
    try:
        with io.BytesIO(xml_content) as xml_stream:
            for event, element in element_tree.iterparse(xml_stream, events=("start", "end")):
                _check_cancel(cancel_check)
                if event == "start" and _local_name(element.tag) == "dimension":
                    reference = element.attrib.get("ref", "")
                    endpoints = reference.split(":")
                    if len(endpoints) not in (1, 2):
                        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                    for endpoint in endpoints:
                        parsed = _parse_cell_reference(endpoint)
                        if parsed is None:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                        column_index, row_number = parsed
                        declared_max_row = max(declared_max_row, row_number)
                        declared_max_column = max(declared_max_column, column_index + 1)
                    if declared_max_row > limits.max_rows or declared_max_column > limits.max_columns:
                        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
    except ImportPreviewError:
        raise
    except element_tree.ParseError:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None

    total_cells = 0
    row_count = 0
    actual_max_row = 0
    header_width: int | None = None
    previous_row = 0
    current_row: int | None = None
    current_max_column = 0
    current_columns: set[int] = set()
    data_row_count = 0

    try:
        with io.BytesIO(xml_content) as xml_stream:
            for event, element in element_tree.iterparse(xml_stream, events=("start", "end")):
                _check_cancel(cancel_check)
                local_name = _local_name(element.tag)
                if event == "start" and local_name == "row":
                    current_row = _positive_xml_integer(element.attrib.get("r"))
                    if current_row is None or current_row <= previous_row:
                        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                    if current_row > limits.max_rows:
                        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                    current_max_column = 0
                    current_columns = set()
                elif event == "end" and local_name == "c":
                    parsed = _parse_cell_reference(element.attrib.get("r"))
                    if parsed is None:
                        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                    column_index, cell_row = parsed
                    if current_row != cell_row or column_index >= limits.max_columns:
                        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                    if column_index in current_columns:
                        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                    current_columns.add(column_index)
                    current_max_column = max(current_max_column, column_index + 1)
                    total_cells += 1
                    if total_cells > limits.max_cells:
                        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                    element.clear()
                elif event == "end" and local_name == "row":
                    if current_row is None:
                        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                    row_count += 1
                    actual_max_row = current_row
                    if current_row == header_row_index + 1:
                        if current_max_column <= 0:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                        header_width = current_max_column
                        if header_width > limits.max_columns:
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                        total_cells += header_width - len(current_columns)
                        if total_cells > limits.max_cells:
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                    elif current_row > header_row_index + 1:
                        if header_width is None:
                            raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
                        gap_start = max(previous_row + 1, header_row_index + 2)
                        gap_rows = max(0, current_row - gap_start)
                        if gap_rows:
                            total_cells += gap_rows * header_width
                            if total_cells > limits.max_cells:
                                raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                            data_row_count += gap_rows
                        if current_max_column > header_width:
                            raise ImportPreviewError(
                                ImportErrorCode.MALFORMED_ROW,
                                row_index=current_row - 1,
                            )
                        total_cells += header_width - len(current_columns)
                        if total_cells > limits.max_cells:
                            raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                        data_row_count += 1
                    previous_row = current_row
                    current_row = None
                    element.clear()
    except ImportPreviewError:
        raise
    except element_tree.ParseError:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT) from None

    if header_width is None or row_count == 0:
        raise ImportPreviewError(ImportErrorCode.INVALID_FORMAT)
    if actual_max_row > limits.max_rows or header_width > limits.max_columns:
        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
    return actual_max_row, header_width, data_row_count


def _fingerprint_path(
    path: Path,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> SourceFingerprint:
    try:
        with path.open("rb") as stream:
            before = os.fstat(stream.fileno())
            if before.st_size > limits.max_input_bytes:
                raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
            digest = hashlib.sha256()
            size = 0
            while True:
                _check_cancel(cancel_check)
                block = stream.read(64 * 1024)
                if not block:
                    break
                size += len(block)
                if size > limits.max_input_bytes:
                    raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                digest.update(block)
            after = os.fstat(stream.fileno())
        path_after = path.stat()
    except ImportPreviewError:
        raise
    except (OSError, ValueError):
        raise ImportPreviewError(ImportErrorCode.INVALID_SOURCE) from None
    if not _same_stat(before, after) or not _same_stat(after, path_after) or size != after.st_size:
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)
    return _fingerprint_from_stat(after, digest.hexdigest())


def _capture_source(
    path: Path,
    suffix: str,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> tuple[SourceFingerprint, _SourceSnapshot]:
    """Hash the complete bounded source while copying the exact bytes to TEMP."""

    temporary_path: Path | None = None
    descriptor: int | None = None
    try:
        with path.open("rb") as source_stream:
            before = os.fstat(source_stream.fileno())
            if before.st_size > limits.max_input_bytes:
                raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
            descriptor, temporary_name = tempfile.mkstemp(prefix="pengtools-import-", suffix=suffix)
            temporary_path = Path(temporary_name)
            digest = hashlib.sha256()
            size = 0
            with os.fdopen(descriptor, "wb") as snapshot_stream:
                descriptor = None
                while True:
                    _check_cancel(cancel_check)
                    block = source_stream.read(64 * 1024)
                    if not block:
                        break
                    size += len(block)
                    if size > limits.max_input_bytes:
                        raise ImportPreviewError(ImportErrorCode.LIMIT_EXCEEDED)
                    snapshot_stream.write(block)
                    digest.update(block)
                snapshot_stream.flush()
                after = os.fstat(source_stream.fileno())
            path_after = path.stat()
        if not _same_stat(before, after) or not _same_stat(after, path_after) or size != after.st_size:
            raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)
        os.chmod(temporary_path, stat.S_IREAD)
        fingerprint = _fingerprint_from_stat(after, digest.hexdigest())
        return fingerprint, _SourceSnapshot(temporary_path)
    except ImportPreviewError:
        if temporary_path is not None:
            _delete_temporary_file(temporary_path)
        raise
    except (OSError, ValueError):
        if temporary_path is not None:
            _delete_temporary_file(temporary_path)
        raise ImportPreviewError(ImportErrorCode.INVALID_SOURCE) from None
    finally:
        if descriptor is not None:
            try:
                os.close(descriptor)
            except OSError:
                pass


def _delete_temporary_file(path: Path) -> None:
    """Remove only a snapshot path created by this module."""

    try:
        os.chmod(path, stat.S_IREAD | stat.S_IWRITE)
    except OSError:
        pass
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def _verify_snapshot(
    path: Path,
    fingerprint: SourceFingerprint,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> None:
    try:
        actual = _fingerprint_path(path, limits, cancel_check)
    except ImportPreviewError as error:
        if error.code in (
            ImportErrorCode.CANCELLED.value,
            ImportErrorCode.CANCEL_CHECK_FAILED.value,
        ):
            raise
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED) from None
    if actual.size_bytes != fingerprint.size_bytes or actual.sha256 != fingerprint.sha256:
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)


def _verify_fingerprint(
    path: Path,
    fingerprint: SourceFingerprint,
    limits: ImportLimits,
    cancel_check: CancelCheck | None,
) -> None:
    try:
        actual = _fingerprint_path(path, limits, cancel_check)
    except ImportPreviewError as error:
        if error.code not in (
            ImportErrorCode.CANCELLED.value,
            ImportErrorCode.CANCEL_CHECK_FAILED.value,
        ):
            raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED) from None
        raise
    if actual != fingerprint:
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)


def _check_path_identity(path: Path, fingerprint: SourceFingerprint) -> None:
    try:
        current = path.stat()
    except OSError:
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED) from None
    if not _stat_matches_fingerprint(current, fingerprint):
        raise ImportPreviewError(ImportErrorCode.SOURCE_CHANGED)


def _same_stat(left: os.stat_result, right: os.stat_result) -> bool:
    return (
        left.st_size == right.st_size
        and left.st_mtime_ns == right.st_mtime_ns
        and left.st_dev == right.st_dev
        and left.st_ino == right.st_ino
    )


def _stat_matches_fingerprint(value: os.stat_result, fingerprint: SourceFingerprint) -> bool:
    return (
        value.st_size == fingerprint.size_bytes
        and value.st_mtime_ns == fingerprint.modified_ns
        and value.st_dev == fingerprint.device
        and value.st_ino == fingerprint.inode
    )


def _fingerprint_from_stat(value: os.stat_result, sha256: str) -> SourceFingerprint:
    return SourceFingerprint(
        size_bytes=value.st_size,
        modified_ns=value.st_mtime_ns,
        device=value.st_dev,
        inode=value.st_ino,
        sha256=sha256,
    )


def _check_cancel(cancel_check: CancelCheck | None) -> None:
    if cancel_check is None:
        return
    try:
        cancelled = cancel_check()
    except Exception:
        raise ImportPreviewError(ImportErrorCode.CANCEL_CHECK_FAILED) from None
    if cancelled is True:
        raise ImportPreviewError(ImportErrorCode.CANCELLED)


def _validate_cell_lengths(
    values: Sequence[str],
    limits: ImportLimits,
    *,
    row_index: int | None = None,
) -> None:
    for column_index, value in enumerate(values):
        if len(value) > limits.max_cell_characters:
            raise ImportPreviewError(
                ImportErrorCode.LIMIT_EXCEEDED,
                row_index=row_index,
                column_index=column_index,
            )


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _positive_xml_integer(value: str | None) -> int | None:
    if value is None or not value.isdigit() or value.startswith("0"):
        return None
    return int(value)


def _parse_cell_reference(value: str | None) -> tuple[int, int] | None:
    if value is None:
        return None
    match = _CELL_REFERENCE_RE.fullmatch(value)
    if match is None:
        return None
    letters, row_text = match.groups()
    column_number = 0
    for character in letters.upper():
        column_number = column_number * 26 + (ord(character) - ord("A") + 1)
    return column_number - 1, int(row_text)


def _safe_zip_member(name: str) -> bool:
    if not name or "\\" in name or name.startswith("/"):
        return False
    parts = name.rstrip("/").split("/")
    return all(part not in ("", ".", "..") for part in parts)


def _resolve_sheet_target(target: str) -> str | None:
    if not isinstance(target, str) or not target or "\\" in target:
        return None
    if target.startswith("/"):
        resolved = posixpath.normpath(target.lstrip("/"))
    else:
        resolved = posixpath.normpath(posixpath.join("xl", target))
    if resolved == "xl" or not resolved.startswith("xl/") or resolved.startswith("../"):
        return None
    return resolved
