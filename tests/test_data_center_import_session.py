from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
import tempfile
import threading
import unittest
from pathlib import Path
import re
from unittest.mock import patch

from openpyxl import Workbook

from tools.data_center.import_preview import (
    ColumnMapping, CsvSource, ImportLimits, ImportValueKind, NullRule, TargetColumn, XlsxSource,
    approve_import, inspect_import_source, preview_import,
)
from tools.data_center.import_session import (
    ImportSession, ImportSessionError, ImportSessionRequest, ImportSessionState,
    candidate_target_columns,
)
from tools.data_center.host_metadata import MetadataTarget
from tools.data_center.manual_session import (
    ManualSession, ManualSessionState, ManualSessionTarget,
    StatementClassification, StatementKind, TransactionEffect,
)
from tools.data_center.manual_relational_providers import BoundManualTableScope, ManualRelationalProviderFactory
from tools.data_center.object_tree import ObjectTreeService, TreeNode, TreeTarget
from tools.data_center.object_tree_composition import ObjectTreeComposition
from tools.data_center.profile_host import ConnectionChoice, DataCenterProfileHost
from tools.data_center.host_relational import RelationalProfileSnapshot, RelationalReadOnlyHost
from tools.data_center.table_source import TableColumnMetadata, TableSourceMetadata


class _DirectoryLease:
    def __init__(self, host, target):
        self.host = host
        self.target = target
        self.dialect = "mysql" if target.provider.endswith("mysql") else "oracle"

    def execute_readonly(self, request, **_kwargs):
        sql = request["sql"].casefold()
        if "information_schema.schemata" in sql:
            rows = [(self.host.database or self.target.database,)]
        elif "information_schema.tables" in sql:
            rows = [(self.host.table_name, "BASE TABLE"), (self.host.view_name, "VIEW")]
        elif "select distinct owner" in sql:
            rows = [(self.host.schema,)]
        elif "select object_name, object_type" in sql:
            rows = [(self.host.table_name, "TABLE"), (self.host.view_name, "VIEW")]
        else:
            rows = []
        offset = int(request["parameters"][-1]) if request["parameters"] else 0
        return {"columns": (), "rows": rows[offset : offset + int(request["row_limit"])]}

    def close(self):
        return None


class _DirectoryFactory:
    def __init__(self, host, target):
        self.host = host
        self.target = target

    def __call__(self, _connection_id):
        return _DirectoryLease(self.host, self.target)


class _DirectoryHost:
    connection_id = "conn-import"
    revision = "rev-1"
    provider = "mysql"
    mode = "mysql"
    database = "sales db"
    table_name = "items"
    view_name = "items_view"

    def profile_snapshot(self, connection_id):
        return type("Snapshot", (), {
            "connection_id": connection_id,
            "profile_revision": self.revision,
            "provider": self.provider,
            "mode": self.mode,
            "database": self.database,
        })()

    def create_factory(self, connection_id, profile_revision, provider, mode, database_name, **_kwargs):
        target = MetadataTarget(connection_id, profile_revision, provider, mode, database_name, ())
        return _DirectoryFactory(self, target)


class _CompositionDirectoryHost(RelationalReadOnlyHost):
    def __init__(self, directory_host):
        self.directory_host = directory_host

    def profile_snapshot(self, connection_id):
        host = self.directory_host
        return RelationalProfileSnapshot(
            connection_id, host.revision, host.provider, host.mode, host.mode,
        )

    def create_factory(
        self, connection_id, *, profile_revision=None, provider=None, mode=None,
        dialect=None, database_name="",
    ):
        return self.directory_host.create_factory(
            connection_id, profile_revision, provider, mode, database_name,
        )


class _Cursor:
    def __init__(self, driver, rows=(), rowcount=0):
        self.driver = driver
        self.rows = list(rows)
        self.rowcount = rowcount
        self.description = None
        self.offset = 0

    def fetchmany(self, size):
        self.driver.events.append(("fetchmany", threading.get_ident()))
        rows = self.rows[self.offset : self.offset + size]
        self.offset += len(rows)
        return rows

    def close(self):
        self.driver.events.append(("cursor_close", threading.get_ident()))


class _ImportDriver:
    def __init__(
        self, *, table_type="BASE TABLE", columns=None, rowcount=None, block_insert=False,
        insert_hook=None, configure_error=False, commit_error=False, rollback_error=False,
        close_failures=0,
    ):
        self.table_type = table_type
        self.mysql_columns = list(columns if columns is not None else [
            ("sales db", "items", "id", "bigint unsigned", "bigint", "NO", 0, "auto_increment", 1),
            ("sales db", "items", "name", "varchar(20)", "varchar", "NO", 0, "", 2),
            ("sales db", "items", "amount", "decimal(6,2)", "decimal", "YES", 0, "", 3),
        ])
        self.rowcount_override = rowcount
        self.block_insert = block_insert
        self.insert_hook = insert_hook
        self.configure_error = configure_error
        self.commit_error = commit_error
        self.rollback_error = rollback_error
        self.close_failures = close_failures
        self.events = []
        self.inserts = []
        self.started_insert = threading.Event()
        self.release_insert = threading.Event()
        self.commit_count = 0
        self.rollback_count = 0
        self.close_count = 0

    @staticmethod
    def _values(parameters):
        if isinstance(parameters, dict):
            return tuple(parameters[key] for key in sorted(parameters))
        return tuple(parameters or ())

    def configure_manual_commit(self):
        self.events.append(("configure", threading.get_ident()))
        if self.configure_error:
            raise RuntimeError("private connector diagnostic")

    def execute(self, sql, parameters):
        values = self._values(parameters)
        self.events.append(("execute", threading.get_ident(), sql))
        lower = sql.casefold()
        if "select table_type from information_schema.tables" in lower:
            return _Cursor(self, [(self.table_type if "items" in values else "VIEW",)])
        if "from information_schema.columns" in lower:
            return _Cursor(self, self.mysql_columns)
        if "from information_schema.key_column_usage" in lower:
            return _Cursor(self, ())
        raise AssertionError("unexpected catalog statement")

    def executemany(self, sql, parameter_rows):
        rows = tuple(tuple(row) for row in parameter_rows)
        self.events.append(("executemany", threading.get_ident(), sql))
        self.inserts.append((sql, rows))
        self.started_insert.set()
        if self.block_insert and not self.release_insert.wait(5):
            raise TimeoutError("fake DBAPI insert wait timed out")
        if self.insert_hook is not None:
            self.insert_hook()
        count = len(rows) if self.rowcount_override is None else self.rowcount_override
        return _Cursor(self, rowcount=count)

    def commit(self):
        self.commit_count += 1
        self.events.append(("commit", threading.get_ident()))
        if self.commit_error:
            raise RuntimeError("private commit diagnostic")

    def rollback(self):
        self.rollback_count += 1
        self.events.append(("rollback", threading.get_ident()))
        if self.rollback_error:
            raise RuntimeError("private rollback diagnostic")

    def close(self):
        self.close_count += 1
        self.events.append(("close", threading.get_ident()))
        if self.close_failures:
            self.close_failures -= 1
            raise RuntimeError("private close diagnostic")


class _ManualOnlyDriver:
    def __init__(self):
        self.commit_count = 0
        self.rollback_count = 0
        self.events = []

    def configure_manual_commit(self):
        self.events.append(("configure", threading.get_ident()))

    def execute(self, sql, _parameters):
        self.events.append(("execute", threading.get_ident(), sql))
        return _Cursor(self, rowcount=1)

    def commit(self):
        self.commit_count += 1

    def rollback(self):
        self.rollback_count += 1

    def close(self):
        return None


class _ProviderCursor:
    def __init__(self, connection):
        self.connection = connection
        self.rows = []
        self.rowcount = -1
        self.description = None
        self.closed = False

    def execute(self, sql, parameters=None):
        self.connection.events.append(("execute", sql, parameters, threading.get_ident()))
        upper = sql.upper()
        if upper.startswith("SELECT SYS_CONTEXT('USERENV', 'CURRENT_SCHEMA')"):
            self.rows = [(self.connection.current_schema,)]
            self.description = (("CURRENT_SCHEMA",),)
            self.rowcount = 1
        elif upper.startswith("ALTER SESSION SET CURRENT_SCHEMA"):
            match = re.fullmatch(
                r'(?is)\s*ALTER\s+SESSION\s+SET\s+CURRENT_SCHEMA\s*=\s*"((?:[^"]|"")*)"\s*',
                sql,
            )
            if not match:
                raise AssertionError("oracle owner name must be identifier-quoted")
            self.connection.current_schema = match.group(1).replace('""', '"')
            self.rows = []
            self.description = ()
            self.rowcount = 0
        elif "FROM ALL_OBJECTS" in upper:
            self.rows = [("TABLE",)]
            self.description = (("OBJECT_TYPE",),)
            self.rowcount = 1
        elif "FROM ALL_TAB_COLUMNS C" in upper:
            self.rows = list(self.connection.oracle_columns)
            self.description = ()
            self.rowcount = len(self.rows)
        elif "FROM ALL_CONSTRAINTS C" in upper:
            self.rows = []
            self.description = ()
            self.rowcount = 0
        else:
            raise AssertionError("unexpected Oracle catalog statement")
        return self

    def executemany(self, sql, rows):
        copied = tuple(dict(row) if isinstance(row, dict) else tuple(row) for row in rows)
        self.connection.events.append(("executemany", sql, copied, threading.get_ident()))
        self.rowcount = len(copied)
        return None

    def fetchone(self):
        return self.rows.pop(0) if self.rows else None

    def fetchmany(self, size):
        rows, self.rows = self.rows[:size], self.rows[size:]
        return rows

    def close(self):
        self.closed = True
        self.connection.events.append(("cursor_close", threading.get_ident()))


class _ProviderConnection:
    def __init__(self, login_schema, table, owner, *, close_failures=0):
        self.current_schema = login_schema
        self.oracle_columns = [
            (owner, table, "name", "VARCHAR2", 80, None, None, "N", 0, "NO", "NO", 1),
            (owner, table, "amount", "NUMBER", 22, 6, 2, "Y", 0, "NO", "NO", 2),
        ]
        self._autocommit = True
        self.events = []
        self.closed = False
        self.close_failures = close_failures
        self.commit_count = 0
        self.rollback_count = 0

    @property
    def autocommit(self):
        return self._autocommit

    @autocommit.setter
    def autocommit(self, value):
        self._autocommit = bool(value)
        self.events.append(("autocommit", bool(value), threading.get_ident()))

    def cursor(self):
        self.events.append(("cursor", threading.get_ident()))
        return _ProviderCursor(self)

    def commit(self):
        self.commit_count += 1
        self.events.append(("commit", threading.get_ident()))

    def rollback(self):
        self.rollback_count += 1
        self.events.append(("rollback", threading.get_ident()))

    def close(self):
        if self.close_failures:
            self.close_failures -= 1
            self.events.append(("connection_close_failed", threading.get_ident()))
            raise RuntimeError("private close diagnostic")
        self.closed = True
        self.events.append(("connection_close", threading.get_ident()))


class _MySqlProviderCursor:
    def __init__(self, connection):
        self.connection = connection
        self.delegate = None
        self.rowcount = -1

    def execute(self, sql, parameters=None):
        self.delegate = self.connection.execute(sql, parameters)
        self.rowcount = self.delegate.rowcount
        return None

    def executemany(self, sql, rows):
        copied = tuple(tuple(row) for row in rows)
        self.connection.events.append(("executemany", threading.get_ident(), sql))
        self.connection.inserts.append((sql, copied))
        self.rowcount = len(copied)
        return None

    def fetchmany(self, size):
        return self.delegate.fetchmany(size) if self.delegate is not None else []

    def fetchone(self):
        rows = self.fetchmany(1)
        return rows[0] if rows else None

    def close(self):
        if self.delegate is not None:
            self.delegate.close()
        else:
            self.connection.events.append(("cursor_close", threading.get_ident()))


class _MySqlProviderConnection(_ImportDriver):
    def __init__(self):
        super().__init__()
        self._autocommit = True
        self.closed = False

    @property
    def autocommit(self):
        return self._autocommit

    @autocommit.setter
    def autocommit(self, value):
        self._autocommit = bool(value)
        self.events.append(("autocommit", bool(value), threading.get_ident()))

    def cursor(self):
        self.events.append(("cursor", threading.get_ident()))
        return _MySqlProviderCursor(self)

    def close(self):
        self.close_count += 1
        self.closed = True
        self.events.append(("connection_close", threading.get_ident()))


class ImportSessionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="pengtools-import-session-")
        self.root = Path(self.temp.name)
        self.host = _DirectoryHost()
        self.tree = ObjectTreeService(relational_host=self.host)
        self.tree_target = TreeTarget(
            self.host.connection_id, self.host.revision, "mysql", mode="mysql", database=self.host.database,
        )
        root = self.tree.root(self.tree_target)
        first = self.tree.submit_children(self.tree_target, root, generation=17).future.result(timeout=3)
        database = next(node for node in first.nodes if node.kind == "database")
        second = self.tree.submit_children(self.tree_target, database, generation=17).future.result(timeout=3)
        self.table = next(node for node in second.nodes if node.kind == "table")
        self.view = next(node for node in second.nodes if node.kind == "view")
        self.choice = ConnectionChoice(
            self.host.connection_id, self.host.revision, "mysql", "mysql", "mysql", self.host.database,
        )
        self.target = ManualSessionTarget(
            "tab-import", self.host.connection_id, self.host.revision,
            database=self.host.database, schema=self.host.database, generation=2,
        )
        self.factory_calls = []

    def tearDown(self):
        self.tree.close().result(timeout=3)
        self.temp.cleanup()

    def _approved(self, text="amount,name\n1.25,Alpha\n2.30,Beta\n", *, limits=None, targets=None, mapping=None):
        path = self.root / "import.csv"
        path.write_text(text, encoding="utf-8", newline="")
        source = CsvSource(path, encoding="utf-8", delimiter=",")
        source_preview = inspect_import_source(source, limits=limits or ImportLimits(max_preview_rows=1))
        selected_targets = tuple(targets or (
            TargetColumn("name", ImportValueKind.TEXT, False),
            TargetColumn("amount", ImportValueKind.DECIMAL, True, 6, 2),
        ))
        selected_mapping = tuple(mapping or (ColumnMapping(1, "name"), ColumnMapping(0, "amount")))
        nulls = NullRule(("NULL",), False, ("NULL",), False, True)
        preview = preview_import(source_preview, selected_mapping, selected_targets, nulls)
        self.assertEqual(preview.issues, ())
        return approve_import(preview, user_approved=True), path

    def _request(self, approved, *, node=None):
        return ImportSessionRequest(
            approved, self.tree, self.tree_target, node or self.table, 17, self.target, self.choice,
        )

    def _session(self, approved, driver, *, node=None):
        def factory(target):
            self.factory_calls.append((threading.get_ident(), target))
            return driver
        return ImportSession(self._request(approved, node=node), driver_factory=factory)

    def _candidate_metadata(self, columns, *, ref=None, is_base_table=True):
        signed_ref = ref or self.tree.resolve_table_source(
            self.tree_target, self.table, generation=17,
        )
        return TableSourceMetadata(
            self.target, signed_ref, tuple(columns), (), is_base_table, None, "",
        )

    def _expect_run_error(self, session, code):
        with self.assertRaises(ImportSessionError) as raised:
            session.start().result(timeout=5)
        self.assertEqual(raised.exception.code, code)
        self.assertEqual(str(raised.exception), code)
        session.cleanup_future.result(timeout=5)

    def test_explicit_approval_insert_and_commit_with_no_primary_key(self):
        approved, _path = self._approved(
            "amount,name\n1.25,雪\nNULL,\n",
            limits=ImportLimits(max_preview_rows=1, max_chunk_rows=1),
        )
        driver = _ImportDriver()
        session = self._session(approved, driver)
        with patch.object(Path, "open", side_effect=AssertionError("constructor must not read files")):
            inert = self._session(approved, _ImportDriver())
            self.assertEqual(inert.status.state, ImportSessionState.CREATED)
        self.assertEqual(self.factory_calls, [])

        run = session.start().result(timeout=5)
        self.assertEqual(run.state, ImportSessionState.READY_TO_COMMIT)
        self.assertEqual((run.rows_processed, run.chunks_processed), (2, 2))
        self.assertEqual((driver.commit_count, driver.rollback_count), (0, 0))
        sql = driver.inserts[0][0]
        tick = chr(96)
        self.assertIn("INSERT INTO " + tick + "sales db" + tick + "." + tick + "items" + tick, sql)
        self.assertNotIn("雪", sql)
        rows = tuple(row for _statement, batch in driver.inserts for row in batch)
        self.assertEqual(rows, (("雪", Decimal("1.25")), ("", None)))
        self.assertFalse(session.cleanup_future.done())

        decision = session.commit().result(timeout=5)
        self.assertTrue(decision.committed)
        self.assertFalse(decision.rolled_back)
        self.assertEqual(session.status.state, ImportSessionState.COMMITTED)
        self.assertEqual((driver.commit_count, driver.rollback_count), (1, 0))
        self.assertTrue(session.cleanup_future.result(timeout=5).complete)
        owner_ids = {event[1] for event in driver.events if len(event) > 1 and isinstance(event[1], int)}
        self.assertEqual(len(owner_ids), 1)
        self.assertNotEqual(next(iter(owner_ids)), threading.get_ident())
        stale_driver = _ImportDriver()
        self._expect_run_error(self._session(approved, stale_driver), "SOURCE_CHANGED")
        self.assertEqual(stale_driver.events, [])

    def test_manual_dirty_transaction_remains_independent(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        manual_driver = _ManualOnlyDriver()
        classify = lambda _sql: StatementClassification(StatementKind.DML, TransactionEffect.DIRTY)
        manual = ManualSession(self.target, lambda _target: manual_driver, classify)
        manual.submit("UPDATE existing SET x = 1").result(timeout=3)
        self.assertEqual(manual.state, ManualSessionState.DIRTY)

        import_driver = _ImportDriver()
        imported = self._session(approved, import_driver)
        imported.start().result(timeout=5)
        imported.commit().result(timeout=5)
        self.assertEqual(manual.state, ManualSessionState.DIRTY)
        self.assertEqual((manual_driver.commit_count, manual_driver.rollback_count), (0, 0))
        self.assertEqual((import_driver.commit_count, import_driver.rollback_count), (1, 0))
        manual.rollback().result(timeout=3)
        manual.close().result(timeout=3)

    def test_empty_mysql_composition_root_uses_resolver_signed_selected_database(self):
        profile = {
            "id": self.host.connection_id,
            "provider": "mysql",
            "engine": "mysql",
            "mode": "mysql",
            "dialect": "mysql",
            "host": "127.0.0.1",
            "port": 3306,
            "username": "TEMP-USER",
            "password": "TEMP-CIPHER",
            "database": self.host.database,
        }
        profile_host = DataCenterProfileHost(connections_loader=lambda: [dict(profile)])
        choice = profile_host.capture_connection(profile, database=self.host.database)
        directory_host = _DirectoryHost()
        directory_host.revision = choice.revision
        composition = ObjectTreeComposition(
            relational_host=_CompositionDirectoryHost(directory_host),
        )
        service = composition.service
        connection = _MySqlProviderConnection()
        connector_profiles = []
        session = None

        def connector(loaded, _secret):
            connector_profiles.append(dict(loaded))
            return connection

        try:
            tree_target = composition.submit_target_capture(
                choice.connection_id, database_scope="",
            ).result(timeout=3)
            self.assertEqual(tree_target.database, "")
            root = service.root(tree_target)
            databases = service.submit_children(
                tree_target, root, generation=41,
            ).future.result(timeout=3)
            database_node = next(node for node in databases.nodes if node.kind == "database")
            self.assertEqual(database_node.database, choice.database)
            tables = service.submit_children(
                tree_target, database_node, generation=41,
            ).future.result(timeout=3)
            table_node = next(node for node in tables.nodes if node.kind == "table")
            self.assertEqual(table_node.database, choice.database)
            target = ManualSessionTarget(
                "tab-global-mysql-import", choice.connection_id, choice.revision,
                database=choice.database, schema=choice.database, generation=8,
            )
            scope = BoundManualTableScope.capture(
                service, tree_target, table_node, 41, target, choice,
            )
            approved, _path = self._approved("amount,name\n1.25,Alpha\n")
            factory = ManualRelationalProviderFactory(
                choice, profile_host=profile_host, connector=connector,
                secret_resolver=lambda _token: "TEMP-SECRET", table_scope=scope,
            )
            request = ImportSessionRequest(
                approved, service, tree_target, table_node, 41, target, choice,
            )
            session = ImportSession(request, driver_factory=factory)
            run = session.start().result(timeout=5)
            self.assertEqual(run.state, ImportSessionState.READY_TO_COMMIT)
            self.assertEqual(connector_profiles[0]["database"], choice.database)
            self.assertEqual(connection.inserts[0][1], (("Alpha", Decimal("1.25")),))
            decision = session.commit().result(timeout=5)
            self.assertTrue(decision.committed)
            self.assertTrue(connection.closed)

            wrong_root = TreeTarget(
                choice.connection_id, choice.revision, "mysql", mode="mysql",
                database="other db",
            )
            with self.assertRaises(ImportSessionError) as nonempty_error:
                ImportSessionRequest(
                    approved, service, wrong_root, table_node, 41, target, choice,
                )
            self.assertEqual(nonempty_error.exception.code, "TARGET_MISMATCH")

            oracle_choice = ConnectionChoice(
                choice.connection_id, choice.revision, "oracle", "oracle", "oracle", "service.prod",
            )
            oracle_root = TreeTarget(
                choice.connection_id, choice.revision, "oracle", mode="oracle", database="",
            )
            oracle_node = TreeNode(
                "signed-oracle-test-node", "table", "items", database="service.prod", schema="OWNER",
            )
            oracle_target = ManualSessionTarget(
                "tab-empty-oracle", choice.connection_id, choice.revision,
                database="service.prod", schema="OWNER", generation=1,
            )
            with self.assertRaises(ImportSessionError) as oracle_error:
                ImportSessionRequest(
                    approved, service, oracle_root, oracle_node, 41, oracle_target, oracle_choice,
                )
            self.assertEqual(oracle_error.exception.code, "TARGET_MISMATCH")
            self.assertEqual(len(connector_profiles), 1)
        finally:
            if session is not None:
                session.close().result(timeout=5)
            composition.close().result(timeout=5)

    def test_candidate_targets_reuse_supported_types_and_allow_base_table_without_key(self):
        metadata = self._candidate_metadata((
            TableColumnMetadata("amount", "decimal(6,2)", True, True, False, False, False, False),
            TableColumnMetadata("name", "varchar(120)", False, True, False, False, False, False),
            TableColumnMetadata("id", "int unsigned", False, None, False, False, False, True),
            TableColumnMetadata("generated_id", "int", False, True, True, False, False, False),
            TableColumnMetadata("payload", "blob", True, True, False, False, False, False),
        ))

        result = candidate_target_columns(metadata)
        by_name = {item.name: item for item in result.columns}

        self.assertTrue(result.can_approve)
        self.assertEqual(result.blocking_reason_codes, ())
        self.assertEqual(by_name["amount"].target, TargetColumn(
            "amount", ImportValueKind.DECIMAL, True, 6, 2,
        ))
        self.assertEqual(by_name["name"].target.kind, ImportValueKind.TEXT)
        self.assertEqual(by_name["id"].target.kind, ImportValueKind.INTEGER)
        self.assertIsNone(by_name["generated_id"].target)
        self.assertEqual(by_name["generated_id"].reason_code, "GENERATED_COLUMN")
        self.assertIsNone(by_name["payload"].target)
        self.assertEqual(by_name["payload"].reason_code, "UNSUPPORTED_TARGET_TYPE")

    def test_unsupported_required_target_blocks_approval_with_actual_reason(self):
        metadata = self._candidate_metadata((
            TableColumnMetadata("name", "varchar(80)", False, True, False, False, False, False),
            TableColumnMetadata("payload", "blob", False, False, False, False, False, True),
        ))

        result = candidate_target_columns(metadata)
        payload = next(item for item in result.columns if item.name == "payload")

        self.assertFalse(result.can_approve)
        self.assertEqual(payload.reason_code, "UNSUPPORTED_TARGET_TYPE")
        self.assertIn("UNSUPPORTED_TARGET_TYPE", result.blocking_reason_codes)

    def test_view_and_unknown_generation_candidates_are_blocked(self):
        view_ref = replace(
            self.tree.resolve_table_source(self.tree_target, self.table, generation=17),
            kind="view",
        )
        view = self._candidate_metadata((
            TableColumnMetadata("name", "varchar(80)", False, True, False, False, False, False),
        ), ref=view_ref, is_base_table=False)
        unknown_generation = self._candidate_metadata((
            TableColumnMetadata("name", "varchar(80)", False, True, None, False, False, False),
        ))

        view_result = candidate_target_columns(view)
        generation_result = candidate_target_columns(unknown_generation)

        self.assertFalse(view_result.can_approve)
        self.assertIn("TABLE_NOT_WRITABLE", view_result.blocking_reason_codes)
        self.assertFalse(generation_result.can_approve)
        self.assertEqual(generation_result.columns[0].reason_code, "GENERATION_METADATA_UNKNOWN")
        self.assertIn("GENERATION_METADATA_UNKNOWN", generation_result.blocking_reason_codes)

    def test_named_xlsx_sheet_reaches_database_as_typed_values(self):
        path = self.root / "append.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Chosen"
        sheet.append(("name", "amount"))
        sheet.append(("Alpha", 1.25))
        sheet.append(("Beta", "NULL"))
        workbook.save(path)
        workbook.close()

        source_preview = inspect_import_source(XlsxSource(path, sheet_name="Chosen"))
        targets = (
            TargetColumn("name", ImportValueKind.TEXT, False),
            TargetColumn("amount", ImportValueKind.DECIMAL, True, 6, 2),
        )
        mapping = (ColumnMapping(0, "name"), ColumnMapping(1, "amount"))
        rules = NullRule(("NULL",), False, ("NULL",), False, True)
        preview = preview_import(source_preview, mapping, targets, rules)
        self.assertEqual(preview.issues, ())
        approved = approve_import(preview, user_approved=True)
        driver = _ImportDriver()
        session = self._session(approved, driver)
        session.start().result(timeout=5)
        session.commit().result(timeout=5)
        self.assertEqual(
            tuple(row for _sql, batch in driver.inserts for row in batch),
            (("Alpha", Decimal("1.25")), ("Beta", None)),
        )

    def test_real_scoped_oracle_factory_and_public_batch_adapter_share_import_owner(self):
        owner = 'Sales "East'
        profile = {
            "id": "oracle-import",
            "provider": "oracle",
            "engine": "oracle",
            "mode": "oracle",
            "dialect": "oracle",
            "host": "127.0.0.1",
            "port": 1521,
            "username": "LOGIN",
            "password": "TEMP-CIPHER",
            "database": "service.prod",
            "schema": "LOGIN",
        }
        profile_host = DataCenterProfileHost(connections_loader=lambda: [dict(profile)])
        choice = profile_host.capture_connection(
            profile,
            database="service.prod",
            schema_allowlist=(owner,),
        )
        directory_host = _DirectoryHost()
        directory_host.connection_id = choice.connection_id
        directory_host.revision = choice.revision
        directory_host.provider = "oracle"
        directory_host.mode = "oracle"
        directory_host.database = "service.prod"
        directory_host.schema = owner
        tree_target = TreeTarget(
            choice.connection_id,
            choice.revision,
            "oracle",
            mode="oracle",
            database="service.prod",
        )
        tree = ObjectTreeService(relational_host=directory_host)
        root = tree.root(tree_target)
        first = tree.submit_children(tree_target, root, generation=23).future.result(timeout=3)
        schema_node = next(node for node in first.nodes if node.kind == "schema")
        second = tree.submit_children(tree_target, schema_node, generation=23).future.result(timeout=3)
        table_node = next(node for node in second.nodes if node.kind == "table")
        target = ManualSessionTarget(
            "tab-oracle-import",
            choice.connection_id,
            choice.revision,
            database="service.prod",
            schema=owner,
            generation=5,
        )
        scope = BoundManualTableScope.capture(
            tree,
            tree_target,
            table_node,
            23,
            target,
            choice,
        )
        approved, _path = self._approved(
            "name,amount\nAlpha,1.25\n",
            mapping=(ColumnMapping(0, "name"), ColumnMapping(1, "amount")),
        )
        connection = _ProviderConnection(profile["schema"], table_node.name, owner)
        connector_profiles = []

        def connector(loaded, _secret):
            connector_profiles.append(dict(loaded))
            return connection

        factory = ManualRelationalProviderFactory(
            choice,
            profile_host=profile_host,
            connector=connector,
            secret_resolver=lambda _token: "TEMP-SECRET",
            table_scope=scope,
        )
        session = ImportSession(
            ImportSessionRequest(approved, tree, tree_target, table_node, 23, target, choice),
            driver_factory=factory,
        )
        rollback_session = None
        try:
            run = session.start().result(timeout=5)
            self.assertEqual(run.state, ImportSessionState.READY_TO_COMMIT)
            self.assertEqual(connection.current_schema, owner)
            self.assertEqual(connector_profiles[0]["schema"], "LOGIN")
            self.assertEqual(profile["schema"], "LOGIN")
            self.assertEqual((connection.commit_count, connection.rollback_count), (0, 0))
            insert = next(event for event in connection.events if event[0] == "executemany")
            self.assertIn('INSERT INTO "Sales ""East"."items"', insert[1])
            self.assertEqual(insert[2], ({"p0": "Alpha", "p1": Decimal("1.25")},))

            owner_ids = {
                event[-1]
                for event in connection.events
                if isinstance(event[-1], int)
            }
            self.assertEqual(len(owner_ids), 1)
            self.assertNotEqual(next(iter(owner_ids)), threading.get_ident())
            tree.invalidate(tree_target, table_node)
            session.commit().result(timeout=5)
            self.assertEqual((connection.commit_count, connection.rollback_count), (1, 0))

            tree.invalidate(tree_target)
            rollback_root = tree.root(tree_target)
            rollback_first = tree.submit_children(
                tree_target, rollback_root, generation=24,
            ).future.result(timeout=3)
            rollback_schema = next(node for node in rollback_first.nodes if node.kind == "schema")
            rollback_second = tree.submit_children(
                tree_target, rollback_schema, generation=24,
            ).future.result(timeout=3)
            rollback_table = next(node for node in rollback_second.nodes if node.kind == "table")
            rollback_target = ManualSessionTarget(
                "tab-oracle-import-rollback", choice.connection_id, choice.revision,
                database="service.prod", schema=owner, generation=6,
            )
            rollback_scope = BoundManualTableScope.capture(
                tree, tree_target, rollback_table, 24, rollback_target, choice,
            )
            rollback_approved, _rollback_path = self._approved(
                "name,amount\nBeta,2.50\n",
                mapping=(ColumnMapping(0, "name"), ColumnMapping(1, "amount")),
            )
            rollback_connection = _ProviderConnection(profile["schema"], rollback_table.name, owner)
            rollback_factory = ManualRelationalProviderFactory(
                choice, profile_host=profile_host,
                connector=lambda _loaded, _secret: rollback_connection,
                secret_resolver=lambda _token: "TEMP-SECRET", table_scope=rollback_scope,
            )
            rollback_session = ImportSession(
                ImportSessionRequest(
                    rollback_approved, tree, tree_target, rollback_table, 24,
                    rollback_target, choice,
                ),
                driver_factory=rollback_factory,
            )
            rollback_run = rollback_session.start().result(timeout=5)
            self.assertEqual(rollback_run.state, ImportSessionState.READY_TO_COMMIT)
            tree.invalidate(tree_target, rollback_table)
            rollback_decision = rollback_session.rollback().result(timeout=5)
            self.assertTrue(rollback_decision.rolled_back)
            self.assertEqual((rollback_connection.commit_count, rollback_connection.rollback_count), (0, 1))
        finally:
            session.close().result(timeout=5)
            if rollback_session is not None:
                rollback_session.close().result(timeout=5)
            tree.close().result(timeout=3)

    def test_real_factory_retains_connector_result_after_scope_changes_and_close_fails(self):
        owner = 'Sales "East'
        profile = {
            "id": "oracle-import", "provider": "oracle", "engine": "oracle",
            "mode": "oracle", "dialect": "oracle", "host": "127.0.0.1",
            "port": 1521, "username": "LOGIN", "password": "TEMP-CIPHER",
            "database": "service.prod", "schema": "LOGIN",
        }
        profile_host = DataCenterProfileHost(connections_loader=lambda: [dict(profile)])
        choice = profile_host.capture_connection(profile, database="service.prod", schema_allowlist=(owner,))
        directory_host = _DirectoryHost()
        directory_host.connection_id = choice.connection_id
        directory_host.revision = choice.revision
        directory_host.provider = "oracle"
        directory_host.mode = "oracle"
        directory_host.database = "service.prod"
        directory_host.schema = owner
        tree_target = TreeTarget(
            choice.connection_id, choice.revision, "oracle", mode="oracle", database="service.prod",
        )
        tree = ObjectTreeService(relational_host=directory_host)
        root = tree.root(tree_target)
        first = tree.submit_children(tree_target, root, generation=23).future.result(timeout=3)
        schema_node = next(node for node in first.nodes if node.kind == "schema")
        second = tree.submit_children(tree_target, schema_node, generation=23).future.result(timeout=3)
        table_node = next(node for node in second.nodes if node.kind == "table")
        target = ManualSessionTarget(
            "tab-oracle-import", choice.connection_id, choice.revision,
            database="service.prod", schema=owner, generation=5,
        )
        scope = BoundManualTableScope.capture(tree, tree_target, table_node, 23, target, choice)
        approved, _path = self._approved(
            "name,amount\nAlpha,1.25\n",
            mapping=(ColumnMapping(0, "name"), ColumnMapping(1, "amount")),
        )
        connection = _ProviderConnection(profile["schema"], table_node.name, owner, close_failures=1)
        connector_profiles = []

        def connector(loaded, _secret):
            connector_profiles.append(dict(loaded))
            tree.invalidate(tree_target, table_node)
            return connection

        factory = ManualRelationalProviderFactory(
            choice, profile_host=profile_host, connector=connector,
            secret_resolver=lambda _token: "TEMP-SECRET", table_scope=scope,
        )
        session = ImportSession(
            ImportSessionRequest(approved, tree, tree_target, table_node, 23, target, choice),
            driver_factory=factory,
        )
        try:
            with self.assertRaises(ImportSessionError):
                session.start().result(timeout=5)
            self.assertEqual(connector_profiles[0]["schema"], "LOGIN")
            self.assertFalse(connection.closed)
            database_calls = {
                "autocommit", "cursor", "execute", "executemany", "commit", "rollback",
            }
            self.assertEqual(
                [event[0] for event in connection.events if event[0] in database_calls],
                [],
            )
            self.assertTrue(session.status.connection_retained)
            self.assertTrue(session.status.owner_alive)
            self.assertFalse(session.cleanup_future.done())
            close_ids = [event[-1] for event in connection.events if event[0].startswith("connection_close")]
            self.assertEqual(len(close_ids), 1)
            self.assertNotEqual(close_ids[0], threading.get_ident())
            owner_thread_id = close_ids[0]

            cleanup = session.retry_cleanup().result(timeout=5)
            self.assertTrue(cleanup.complete)
            self.assertFalse(cleanup.connection_retained)
            self.assertTrue(connection.closed)
            close_ids = [event[-1] for event in connection.events if event[0] == "connection_close"]
            self.assertEqual(len(close_ids), 1)
            all_close_ids = {
                event[-1] for event in connection.events
                if event[0].startswith("connection_close")
            }
            self.assertEqual(all_close_ids, {owner_thread_id})
            self.assertNotEqual(owner_thread_id, threading.get_ident())
        finally:
            session.close().result(timeout=5)
            tree.close().result(timeout=3)

    def test_required_unmapped_column_and_view_are_rejected(self):
        approved, _path = self._approved(
            "amount\n1.25\n",
            targets=(TargetColumn("amount", ImportValueKind.DECIMAL, True, 6, 2),),
            mapping=(ColumnMapping(0, "amount"),),
        )
        driver = _ImportDriver()
        self._expect_run_error(self._session(approved, driver), "REQUIRED_COLUMN_UNMAPPED")
        self.assertEqual((driver.inserts, driver.rollback_count), ([], 1))

        approved_view, _path = self._approved("amount,name\n1.25,Alpha\n")
        view_columns = tuple(tuple((row[0], "items_view", *row[2:])) for row in _ImportDriver().mysql_columns)
        view_driver = _ImportDriver(table_type="VIEW", columns=view_columns)
        self._expect_run_error(self._session(approved_view, view_driver, node=self.view), "TABLE_NOT_WRITABLE")
        self.assertEqual((view_driver.inserts, view_driver.rollback_count), ([], 1))

    def test_unknown_rowcount_and_late_conversion_failure_rollback(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        driver = _ImportDriver(rowcount=-1)
        self._expect_run_error(self._session(approved, driver), "ROWCOUNT_UNKNOWN")
        self.assertEqual((driver.commit_count, driver.rollback_count), (0, 1))

        limits = ImportLimits(max_preview_rows=1, max_chunk_rows=1)
        late_bad, _path = self._approved("amount,name\n1.25,Alpha\n99999.99,Beta\n", limits=limits)
        late_driver = _ImportDriver()
        self._expect_run_error(self._session(late_bad, late_driver), "CONVERSION_ERROR")
        self.assertEqual(len(late_driver.inserts), 1)
        self.assertEqual((late_driver.commit_count, late_driver.rollback_count), (0, 1))

    def test_autocommit_setup_failure_rolls_back_and_insert_failure_unknown_is_preserved(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        setup_driver = _ImportDriver(configure_error=True)
        setup_session = self._session(approved, setup_driver)
        self._expect_run_error(setup_session, "AUTOCOMMIT_CONTROL_UNVERIFIED")
        self.assertEqual((setup_driver.inserts, setup_driver.rollback_count, setup_driver.close_count), ([], 1, 1))
        self.assertEqual(setup_session.status.transaction_outcome, "rolled_back")

        approved2, _path = self._approved("amount,name\n1.25,Alpha\n")
        uncertain_driver = _ImportDriver(rowcount=-1, rollback_error=True)
        uncertain_session = self._session(approved2, uncertain_driver)
        self._expect_run_error(uncertain_session, "ROLLBACK_OUTCOME_UNKNOWN")
        self.assertEqual(uncertain_session.status.state, ImportSessionState.UNKNOWN)
        self.assertEqual(uncertain_session.status.transaction_outcome, "unknown")

    def test_changed_original_after_final_partial_chunk_rolls_back(self):
        approved, path = self._approved("amount,name\n1.25,Alpha\n")
        def alter_source():
            path.write_bytes(b"amount,name\n9.99,Other\n")
        driver = _ImportDriver(insert_hook=alter_source)
        session = self._session(approved, driver)
        self._expect_run_error(session, "SOURCE_CHANGED")
        self.assertEqual(session.status.state, ImportSessionState.FAILED)
        self.assertEqual(session.status.transaction_outcome, "rolled_back")
        self.assertEqual(driver.inserts[0][1], (("Alpha", Decimal("1.25")),))
        self.assertEqual((driver.commit_count, driver.rollback_count), (0, 1))

    def test_running_cancel_waits_for_blocked_insert_then_rolls_back(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        driver = _ImportDriver(block_insert=True)
        session = self._session(approved, driver)
        run = session.start()
        self.assertTrue(driver.started_insert.wait(3))
        self.assertFalse(run.cancel())
        self.assertTrue(session.cancel_current())
        self.assertEqual(session.status.state, ImportSessionState.CANCEL_PENDING)
        self.assertFalse(session.cleanup_future.done())
        self.assertEqual(driver.rollback_count, 0)
        driver.release_insert.set()
        with self.assertRaises(ImportSessionError) as raised:
            run.result(timeout=5)
        self.assertEqual(raised.exception.code, "CANCELLED")
        self.assertEqual(session.status.state, ImportSessionState.ROLLED_BACK)
        self.assertEqual((driver.commit_count, driver.rollback_count), (0, 1))

    def test_cancel_between_final_check_and_ready_state_cannot_be_committed(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        driver = _ImportDriver()
        session = self._session(approved, driver)
        original_check = session._check_cancel
        checks = 0

        def cancel_after_final_check():
            nonlocal checks
            original_check()
            checks += 1
            if checks == 6:
                self.assertTrue(session.cancel_current())

        with patch.object(session, "_check_cancel", side_effect=cancel_after_final_check):
            self._expect_run_error(session, "CANCELLED")

        self.assertEqual(session.status.state, ImportSessionState.ROLLED_BACK)
        self.assertEqual((driver.commit_count, driver.rollback_count), (0, 1))

    def test_queued_cancel_uses_no_database_and_releases_snapshot(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        driver = _ImportDriver()
        session = self._session(approved, driver)
        with patch("tools.data_center.import_session.threading.Thread.start"):
            run = session.start()
        self.assertEqual(session.status.state, ImportSessionState.QUEUED)
        self.assertTrue(session.cancel_current())
        self.assertTrue(run.cancelled())
        session._owner.run()
        self.assertEqual(session.status.state, ImportSessionState.CANCELLED_BEFORE_START)
        self.assertEqual(driver.events, [])
        self.assertTrue(session.cleanup_future.result(timeout=3).complete)

    def test_commit_and_rollback_unknown_are_not_reported_as_success(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        commit_driver = _ImportDriver(commit_error=True)
        commit_session = self._session(approved, commit_driver)
        commit_session.start().result(timeout=5)
        with self.assertRaises(ImportSessionError) as commit_error:
            commit_session.commit().result(timeout=5)
        self.assertEqual(commit_error.exception.code, "COMMIT_OUTCOME_UNKNOWN")
        self.assertEqual(commit_session.status.state, ImportSessionState.UNKNOWN)
        self.assertEqual(commit_session.status.transaction_outcome, "unknown")

        approved2, _path = self._approved("amount,name\n1.25,Alpha\n")
        rollback_driver = _ImportDriver(rollback_error=True)
        rollback_session = self._session(approved2, rollback_driver)
        rollback_session.start().result(timeout=5)
        with self.assertRaises(ImportSessionError) as rollback_error:
            rollback_session.rollback().result(timeout=5)
        self.assertEqual(rollback_error.exception.code, "ROLLBACK_OUTCOME_UNKNOWN")
        self.assertEqual(rollback_session.status.state, ImportSessionState.UNKNOWN)
        self.assertEqual(rollback_session.status.transaction_outcome, "unknown")

    def test_failed_close_retains_connection_until_owner_retry(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        driver = _ImportDriver(close_failures=1)
        session = self._session(approved, driver)
        session.start().result(timeout=5)
        decision = session.commit().result(timeout=5)
        self.assertTrue(decision.committed)
        self.assertFalse(decision.cleanup_complete)
        self.assertTrue(session.status.owner_alive)
        self.assertTrue(session.status.connection_retained)
        self.assertFalse(session.cleanup_future.done())
        self.assertTrue(session.retry_cleanup().result(timeout=5).complete)
        self.assertFalse(session.status.connection_retained)
        self.assertTrue(session.cleanup_future.result(timeout=5).complete)
        self.assertEqual(driver.close_count, 2)

    def test_cancel_before_start_closes_snapshot_without_factory_call(self):
        approved, _path = self._approved("amount,name\n1.25,Alpha\n")
        session = ImportSession(
            self._request(approved), driver_factory=lambda _target: self.factory_calls.append(1),
        )
        self.assertTrue(session.close().result(timeout=3).complete)
        self.assertEqual(session.status.state, ImportSessionState.CANCELLED_BEFORE_START)
        self.assertEqual(self.factory_calls, [])


if __name__ == "__main__":
    unittest.main()
