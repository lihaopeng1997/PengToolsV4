from __future__ import annotations

import csv
import datetime
import os
import re
import tempfile
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from tools.data_center import import_preview as import_preview_module
from tools.data_center.import_preview import (
    ApprovedImport,
    ColumnMapping,
    CsvSource,
    ImportErrorCode,
    ImportLimits,
    ImportPreviewError,
    ImportValueKind,
    NullRule,
    RawCellKind,
    TargetColumn,
    XlsxSource,
    approve_import,
    close_import_snapshot,
    inspect_import_source,
    iter_import_chunks,
    preview_import,
)


class ImportPreviewTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def _csv_file(self, name: str, rows, *, encoding: str = "utf-8-sig") -> Path:
        path = self.root / name
        with path.open("w", encoding=encoding, newline="") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            writer.writerows(rows)
        return path

    def _xlsx_file(self, name: str, rows, *, sheet_name: str = "Data") -> Path:
        path = self.root / name
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    @staticmethod
    def _rules() -> NullRule:
        return NullRule(
            csv_null_tokens=(r"\N",),
            csv_empty_is_null=False,
            xlsx_null_tokens=(r"\N",),
            xlsx_empty_is_null=False,
            xlsx_blank_is_null=True,
        )

    @staticmethod
    def _targets() -> tuple[TargetColumn, ...]:
        return (
            TargetColumn("left_name", ImportValueKind.TEXT, nullable=False),
            TargetColumn("right_name", ImportValueKind.TEXT, nullable=False),
            TargetColumn("amount", ImportValueKind.DECIMAL, nullable=False, decimal_precision=22, decimal_scale=2),
            TargetColumn("created_on", ImportValueKind.DATE, nullable=False),
            TargetColumn("empty_text", ImportValueKind.TEXT, nullable=True),
            TargetColumn("optional_text", ImportValueKind.TEXT, nullable=True),
        )

    @staticmethod
    def _mappings() -> tuple[ColumnMapping, ...]:
        return (
            ColumnMapping(0, "left_name"),
            ColumnMapping(1, "right_name"),
            ColumnMapping(2, "amount"),
            ColumnMapping(3, "created_on", "%Y-%m-%d"),
            ColumnMapping(4, "empty_text"),
            ColumnMapping(5, "optional_text"),
        )

    def _approve(self, source, *, limits=ImportLimits(), mapping=None, targets=None, rules=None) -> ApprovedImport:
        source_preview = inspect_import_source(source, limits=limits)
        import_preview = preview_import(
            source_preview,
            self._mappings() if mapping is None else mapping,
            self._targets() if targets is None else targets,
            self._rules() if rules is None else rules,
        )
        return approve_import(import_preview, user_approved=True)

    def test_csv_bom_duplicate_headers_null_empty_decimal_and_explicit_date(self) -> None:
        path = self._csv_file(
            "typed.csv",
            (
                ("名称", "名称", "amount", "created", "empty", "optional"),
                ("左", "右", "12345678901234567890.12", "2026-10-04", "", r"\N"),
                ("甲", "乙", "1.20", "2026-10-05", "x", "present"),
            ),
        )
        limits = ImportLimits(max_preview_rows=1, max_chunk_rows=1)
        source = CsvSource(path, encoding="utf-8-sig", delimiter=",")

        inspected = inspect_import_source(source, limits=limits)
        self.assertEqual(tuple(column.label for column in inspected.columns), ("名称", "名称", "amount", "created", "empty", "optional"))
        self.assertTrue(inspected.has_more)
        preview = preview_import(inspected, self._mappings(), self._targets(), self._rules())
        self.assertEqual(preview.issues, ())
        self.assertEqual(
            preview.rows[0].values,
            ("左", "右", Decimal("12345678901234567890.12"), datetime.date(2026, 10, 4), "", None),
        )

        approved = approve_import(preview, user_approved=True)
        chunks = tuple(iter_import_chunks(approved, chunk_rows=1))
        self.assertEqual([chunk.rows[0][0:2] for chunk in chunks], [("左", "右"), ("甲", "乙")])
        self.assertEqual(chunks[0].rows[0][2], Decimal("12345678901234567890.12"))

    def test_gb18030_is_used_only_when_explicitly_selected(self) -> None:
        path = self._csv_file("gb18030.csv", (("名称",), ("中文",)), encoding="gb18030")
        source = CsvSource(path, encoding="gb18030", delimiter=",")
        inspected = inspect_import_source(source)
        self.assertEqual(inspected.columns[0].label, "名称")
        self.assertEqual(inspected.rows[0].cells[0].value, "中文")

        wrong_encoding = CsvSource(path, encoding="utf-8", delimiter=",")
        with self.assertRaises(ImportPreviewError) as raised:
            inspect_import_source(wrong_encoding)
        self.assertEqual(raised.exception.code, ImportErrorCode.INVALID_ENCODING.value)
        self.assertNotIn(str(path), str(raised.exception))

    def test_source_construction_does_not_read_the_file(self) -> None:
        missing = self.root / "does-not-exist.csv"
        source = CsvSource(missing, encoding="utf-8", delimiter=",")
        self.assertEqual(source.path, missing)
        self.assertFalse(missing.exists())

    def test_mapping_is_index_addressed_and_target_names_must_be_unique(self) -> None:
        path = self._csv_file("duplicates.csv", (("same", "same"), ("first", "second")))
        inspected = inspect_import_source(CsvSource(path, encoding="utf-8-sig", delimiter=","))
        targets = (
            TargetColumn("first_target", ImportValueKind.TEXT, nullable=False),
            TargetColumn("second_target", ImportValueKind.TEXT, nullable=False),
        )
        mapped = preview_import(
            inspected,
            (ColumnMapping(1, "second_target"), ColumnMapping(0, "first_target")),
            targets,
            self._rules(),
        )
        self.assertEqual(mapped.rows[0].values, ("second", "first"))
        with self.assertRaises(ImportPreviewError) as raised:
            preview_import(
                inspected,
                (ColumnMapping(0, "same"), ColumnMapping(1, "same")),
                (
                    TargetColumn("same", ImportValueKind.TEXT, nullable=False),
                    TargetColumn("same", ImportValueKind.TEXT, nullable=False),
                ),
                self._rules(),
            )
        self.assertEqual(raised.exception.code, ImportErrorCode.INVALID_MAPPING.value)

    def test_csv_missing_and_extra_fields_are_rejected(self) -> None:
        for name, content in (
            ("missing.csv", "a,b\none\n"),
            ("extra.csv", "a,b\none,two,three\n"),
        ):
            with self.subTest(name=name):
                path = self.root / name
                path.write_text(content, encoding="utf-8")
                with self.assertRaises(ImportPreviewError) as raised:
                    inspect_import_source(CsvSource(path, encoding="utf-8", delimiter=","))
                self.assertEqual(raised.exception.code, ImportErrorCode.MALFORMED_ROW.value)

    def test_decimal_scale_rejects_rounding_and_approval_rejects_preview_issues(self) -> None:
        path = self._csv_file("scale.csv", (("amount",), ("1.234",)))
        source_preview = inspect_import_source(CsvSource(path, encoding="utf-8-sig", delimiter=","))
        preview = preview_import(
            source_preview,
            (ColumnMapping(0, "amount"),),
            (TargetColumn("amount", ImportValueKind.DECIMAL, nullable=False, decimal_precision=6, decimal_scale=2),),
            self._rules(),
        )
        self.assertEqual(preview.issues[0].code, ImportErrorCode.CONVERSION_ERROR.value)
        with self.assertRaises(ImportPreviewError) as raised:
            approve_import(preview, user_approved=True)
        self.assertEqual(raised.exception.code, ImportErrorCode.PREVIEW_HAS_ERRORS.value)

    def test_approval_must_be_an_explicit_user_action(self) -> None:
        path = self._csv_file("approval.csv", (("value",), ("ok",)))
        source_preview = inspect_import_source(CsvSource(path, encoding="utf-8-sig", delimiter=","))
        preview = preview_import(
            source_preview,
            (ColumnMapping(0, "value"),),
            (TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
            self._rules(),
        )
        with self.assertRaises(ImportPreviewError) as raised:
            approve_import(preview, user_approved=False)
        self.assertEqual(raised.exception.code, ImportErrorCode.APPROVAL_REQUIRED.value)

    def test_cancellation_and_file_limit_have_safe_errors(self) -> None:
        path = self._csv_file("bounded.csv", (("value",), ("some long value",)))
        with self.assertRaises(ImportPreviewError) as cancelled:
            inspect_import_source(
                CsvSource(path, encoding="utf-8-sig", delimiter=","),
                cancel_check=lambda: True,
            )
        self.assertEqual(cancelled.exception.code, ImportErrorCode.CANCELLED.value)
        self.assertNotIn(str(path), str(cancelled.exception))

        with self.assertRaises(ImportPreviewError) as limited:
            inspect_import_source(
                CsvSource(path, encoding="utf-8-sig", delimiter=","),
                limits=ImportLimits(max_input_bytes=4),
            )
        self.assertEqual(limited.exception.code, ImportErrorCode.LIMIT_EXCEEDED.value)

    def test_preview_is_finite_and_chunks_are_lazy(self) -> None:
        path = self._csv_file(
            "lazy.csv",
            (("amount",), ("1.00",), ("not-a-number",), ("3.00",)),
        )
        limits = ImportLimits(max_preview_rows=1, max_chunk_rows=1)
        source_preview = inspect_import_source(CsvSource(path, encoding="utf-8-sig", delimiter=","), limits=limits)
        self.assertEqual(len(source_preview.rows), 1)
        self.assertTrue(source_preview.has_more)
        preview = preview_import(
            source_preview,
            (ColumnMapping(0, "amount"),),
            (TargetColumn("amount", ImportValueKind.DECIMAL, nullable=False, decimal_precision=8, decimal_scale=2),),
            self._rules(),
        )
        approved = approve_import(preview, user_approved=True)
        chunks = iter_import_chunks(approved, chunk_rows=1)
        self.assertEqual(next(chunks).rows, ((Decimal("1.00"),),))
        with self.assertRaises(ImportPreviewError) as raised:
            next(chunks)
        self.assertEqual(raised.exception.code, ImportErrorCode.CONVERSION_ERROR.value)
        self.assertEqual(raised.exception.row_index, 2)

    def test_fingerprint_rechecked_before_first_chunk(self) -> None:
        path = self._csv_file("stale.csv", (("value",), ("one",)))
        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        before = path.stat()
        path.write_text("value\ntwo\n", encoding="utf-8-sig")
        os.utime(path, ns=(before.st_atime_ns, before.st_mtime_ns))
        with self.assertRaises(ImportPreviewError) as raised:
            next(iter_import_chunks(approved, chunk_rows=1))
        self.assertEqual(raised.exception.code, ImportErrorCode.SOURCE_CHANGED.value)

    def test_csv_transient_same_size_change_uses_approved_snapshot(self) -> None:
        path = self._csv_file("transient.csv", (("value",), ("approved",)))
        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        snapshot_path = approved._snapshot.path
        original_bytes = path.read_bytes()
        original_stat = path.stat()
        tampered_bytes = original_bytes.replace(b"approved", b"tampered", 1)
        self.assertEqual(len(tampered_bytes), len(original_bytes))
        real_reader = import_preview_module._iter_csv_data_rows

        def transient_reader(source, columns, limits, cancel_check):
            self.assertEqual(source.path, snapshot_path)
            self.assertNotEqual(source.path, path)
            path.write_bytes(tampered_bytes)
            os.utime(path, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))
            try:
                yield from real_reader(source, columns, limits, cancel_check)
            finally:
                path.write_bytes(original_bytes)
                os.utime(path, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))

        with patch("tools.data_center.import_preview._iter_csv_data_rows", transient_reader):
            chunks = tuple(iter_import_chunks(approved, chunk_rows=1))
        self.assertEqual(tuple(row for chunk in chunks for row in chunk.rows), (("approved",),))
        self.assertEqual(path.read_bytes(), original_bytes)
        self.assertFalse(snapshot_path.exists())

    def test_xlsx_transient_same_size_change_uses_approved_snapshot(self) -> None:
        path = self._xlsx_file("transient.xlsx", (("value",), ("approved",)))
        self._repack_xlsx_stored(path)
        approved = self._approve(
            XlsxSource(path, sheet_name="Data"),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        snapshot_path = approved._snapshot.path
        original_bytes = path.read_bytes()
        self.assertIn(b"approved", original_bytes)
        self.assertNotIn(b"tampered", original_bytes)
        tampered_bytes = original_bytes.replace(b"approved", b"tampered", 1)
        self.assertEqual(len(tampered_bytes), len(original_bytes))
        original_stat = path.stat()
        real_reader = import_preview_module._iter_xlsx_data_rows

        def transient_reader(source, columns, limits, metadata, cancel_check):
            self.assertEqual(source.path, snapshot_path)
            self.assertNotEqual(source.path, path)
            path.write_bytes(tampered_bytes)
            os.utime(path, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))
            try:
                yield from real_reader(source, columns, limits, metadata, cancel_check)
            finally:
                path.write_bytes(original_bytes)
                os.utime(path, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))

        with patch("tools.data_center.import_preview._iter_xlsx_data_rows", transient_reader):
            chunks = tuple(iter_import_chunks(approved, chunk_rows=1))
        self.assertEqual(tuple(row for chunk in chunks for row in chunk.rows), (("approved",),))
        self.assertEqual(path.read_bytes(), original_bytes)
        self.assertFalse(snapshot_path.exists())

    def test_snapshot_is_one_use_and_cancel_before_yield_cleans_up(self) -> None:
        path = self._csv_file("one-use.csv", (("value",), ("one",)))
        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        snapshot_path = approved._snapshot.path
        self.assertIn(Path(tempfile.gettempdir()).resolve(), snapshot_path.resolve().parents)
        chunks = iter_import_chunks(approved, chunk_rows=1, cancel_check=lambda: True)
        with self.assertRaises(ImportPreviewError) as cancelled:
            next(chunks)
        self.assertEqual(cancelled.exception.code, ImportErrorCode.CANCELLED.value)
        self.assertFalse(snapshot_path.exists())
        with self.assertRaises(ImportPreviewError) as reused:
            next(iter_import_chunks(approved, chunk_rows=1))
        self.assertEqual(reused.exception.code, ImportErrorCode.SOURCE_CHANGED.value)

    def test_explicit_snapshot_close_waits_for_active_iterator_and_abandoned_preview_cleans_up(self) -> None:
        path = self._csv_file("explicit-close.csv", (("value",), ("one",), ("two",)))
        source_preview = inspect_import_source(CsvSource(path, encoding="utf-8-sig", delimiter=","))
        abandoned_snapshot = source_preview._snapshot.path
        close_import_snapshot(source_preview)
        self.assertFalse(abandoned_snapshot.exists())

        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            limits=ImportLimits(max_chunk_rows=1),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        snapshot_path = approved._snapshot.path
        chunks = iter_import_chunks(approved, chunk_rows=1)
        self.assertEqual(next(chunks).rows, (("one",),))
        close_import_snapshot(approved)
        self.assertTrue(snapshot_path.exists())
        chunks.close()
        self.assertFalse(snapshot_path.exists())

    def test_xlsx_cancellation_after_yield_cleans_snapshot(self) -> None:
        path = self._xlsx_file("cancel-xlsx.xlsx", (("value",), ("one",), ("two",)))
        approved = self._approve(
            XlsxSource(path, sheet_name="Data"),
            limits=ImportLimits(max_chunk_rows=1),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        snapshot_path = approved._snapshot.path
        state = {"cancelled": False}
        chunks = iter_import_chunks(approved, chunk_rows=1, cancel_check=lambda: state["cancelled"])
        self.assertEqual(next(chunks).rows, (("one",),))
        state["cancelled"] = True
        with self.assertRaises(ImportPreviewError) as cancelled:
            next(chunks)
        self.assertEqual(cancelled.exception.code, ImportErrorCode.CANCELLED.value)
        self.assertFalse(snapshot_path.exists())

    def test_file_change_after_a_yield_requires_consumer_rollback(self) -> None:
        path = self._csv_file("midstream.csv", (("value",), ("one",), ("two",)))
        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            limits=ImportLimits(max_chunk_rows=1),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        chunks = iter_import_chunks(approved, chunk_rows=1)
        self.assertEqual(next(chunks).rows, (("one",),))
        before = path.stat()
        content = path.read_bytes().replace(b"two", b"six")
        path.write_bytes(content)
        os.utime(path, ns=(before.st_atime_ns, before.st_mtime_ns))
        self.assertEqual(next(chunks).rows, (("two",),))
        with self.assertRaises(ImportPreviewError) as raised:
            next(chunks)
        self.assertEqual(raised.exception.code, ImportErrorCode.SOURCE_CHANGED.value)

    def test_cancellation_after_a_yield_stops_the_lazy_iterator(self) -> None:
        path = self._csv_file("cancel-chunk.csv", (("value",), ("one",), ("two",)))
        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            limits=ImportLimits(max_chunk_rows=1),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        state = {"cancelled": False}
        chunks = iter_import_chunks(approved, chunk_rows=1, cancel_check=lambda: state["cancelled"])
        self.assertEqual(next(chunks).rows, (("one",),))
        state["cancelled"] = True
        with self.assertRaises(ImportPreviewError) as raised:
            next(chunks)
        self.assertEqual(raised.exception.code, ImportErrorCode.CANCELLED.value)

    def test_cancellation_after_final_partial_chunk_is_not_successful_exhaustion(self) -> None:
        path = self._csv_file("cancel-final-partial.csv", (("value",), ("one",)))
        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        snapshot_path = approved._snapshot.path
        state = {"cancelled": False}
        chunks = iter_import_chunks(approved, chunk_rows=10, cancel_check=lambda: state["cancelled"])
        self.assertEqual(next(chunks).rows, (("one",),))
        state["cancelled"] = True
        with self.assertRaises(ImportPreviewError) as cancelled:
            next(chunks)
        self.assertEqual(cancelled.exception.code, ImportErrorCode.CANCELLED.value)
        self.assertFalse(snapshot_path.exists())

    def test_source_change_after_final_partial_chunk_is_not_successful_exhaustion(self) -> None:
        path = self._csv_file("source-final-partial.csv", (("value",), ("one",)))
        approved = self._approve(
            CsvSource(path, encoding="utf-8-sig", delimiter=","),
            mapping=(ColumnMapping(0, "value"),),
            targets=(TargetColumn("value", ImportValueKind.TEXT, nullable=False),),
        )
        snapshot_path = approved._snapshot.path
        chunks = iter_import_chunks(approved, chunk_rows=10)
        self.assertEqual(next(chunks).rows, (("one",),))
        before = path.stat()
        path.write_bytes(path.read_bytes().replace(b"one", b"two"))
        os.utime(path, ns=(before.st_atime_ns, before.st_mtime_ns))
        with self.assertRaises(ImportPreviewError) as changed:
            next(chunks)
        self.assertEqual(changed.exception.code, ImportErrorCode.SOURCE_CHANGED.value)
        self.assertFalse(snapshot_path.exists())

    def test_xlsx_preserves_raw_decimal_lexical_date_and_blank(self) -> None:
        path = self._xlsx_file(
            "typed.xlsx",
            (("amount", "created_on", "empty_text"), (0.125, datetime.date(2026, 10, 4), "")),
        )
        self._replace_xlsx_number(path, "A2", "0.123456789012345678901234")
        source = XlsxSource(path, sheet_name="Data")
        inspected = inspect_import_source(source)
        preview = preview_import(
            inspected,
            (
                ColumnMapping(0, "amount"),
                ColumnMapping(1, "created_on"),
                ColumnMapping(2, "empty_text"),
            ),
            (
                TargetColumn("amount", ImportValueKind.DECIMAL, nullable=False, decimal_precision=24, decimal_scale=24),
                TargetColumn("created_on", ImportValueKind.DATE, nullable=False),
                TargetColumn("empty_text", ImportValueKind.TEXT, nullable=True),
            ),
            self._rules(),
        )
        self.assertEqual(preview.issues, ())
        self.assertEqual(
            preview.rows[0].values,
            (Decimal("0.123456789012345678901234"), datetime.date(2026, 10, 4), None),
        )
        approved = approve_import(preview, user_approved=True)
        chunk = next(iter_import_chunks(approved, chunk_rows=1))
        self.assertEqual(chunk.rows, (preview.rows[0].values,))

    def test_xlsx_omitted_trailing_cells_and_all_blank_rows_are_explicit_blanks(self) -> None:
        path = self._xlsx_file(
            "trailing-blanks.xlsx",
            (
                ("first", "second", "optional"),
                ("one", "two"),
                (None, None, None),
                ("three", "four", "five"),
            ),
        )
        limits = ImportLimits(max_preview_rows=10, max_chunk_rows=2)
        inspected = inspect_import_source(
            XlsxSource(path, sheet_name="Data"), limits=limits,
        )
        self.assertEqual(tuple(column.label for column in inspected.columns), ("first", "second", "optional"))
        self.assertEqual(tuple(len(row.cells) for row in inspected.rows), (3, 3, 3))
        self.assertEqual(inspected.rows[0].cells[2].kind, RawCellKind.BLANK)
        self.assertEqual(tuple(cell.kind for cell in inspected.rows[1].cells), (RawCellKind.BLANK,) * 3)
        self.assertEqual(tuple(cell.value for cell in inspected.rows[2].cells), ("three", "four", "five"))
        with self.assertRaises(ImportPreviewError) as cell_budget:
            inspect_import_source(
                XlsxSource(path, sheet_name="Data"),
                limits=ImportLimits(max_preview_rows=10, max_cells=8),
            )
        self.assertEqual(cell_budget.exception.code, ImportErrorCode.LIMIT_EXCEEDED.value)

        mapping = (
            ColumnMapping(0, "first"),
            ColumnMapping(1, "second"),
            ColumnMapping(2, "optional"),
        )
        targets = (
            TargetColumn("first", ImportValueKind.TEXT, nullable=True),
            TargetColumn("second", ImportValueKind.TEXT, nullable=True),
            TargetColumn("optional", ImportValueKind.TEXT, nullable=True),
        )
        preview = preview_import(inspected, mapping, targets, self._rules())
        self.assertEqual(preview.issues, ())
        self.assertEqual(preview.rows[0].values, ("one", "two", None))
        self.assertEqual(preview.rows[1].values, (None, None, None))
        self.assertEqual(preview.rows[2].values, ("three", "four", "five"))

        approved = approve_import(preview, user_approved=True)
        snapshot_path = approved._snapshot.path
        chunks = tuple(iter_import_chunks(approved, chunk_rows=2))
        self.assertEqual(tuple(row for chunk in chunks for row in chunk.rows), tuple(row.values for row in preview.rows))
        self.assertFalse(snapshot_path.exists())

    def test_xlsx_source_requires_named_sheet_and_formula_is_never_evaluated(self) -> None:
        path = self._xlsx_file("sheet.xlsx", (("count",), (1,)), sheet_name="Chosen")
        source = XlsxSource(path, sheet_name="Other")
        with self.assertRaises(ImportPreviewError) as raised:
            inspect_import_source(source)
        self.assertEqual(raised.exception.code, ImportErrorCode.INVALID_SHEET.value)

        selected = XlsxSource(path, sheet_name="Chosen")
        real_load = __import__("openpyxl").load_workbook
        with patch("openpyxl.load_workbook", wraps=real_load) as loader:
            inspect_import_source(selected)
        loader.assert_called_once()
        self.assertTrue(loader.call_args.kwargs["read_only"])
        self.assertFalse(loader.call_args.kwargs["data_only"])
        self.assertFalse(loader.call_args.kwargs["keep_links"])

    def test_xlsx_duplicate_headers_map_by_physical_column_index(self) -> None:
        path = self._xlsx_file("duplicate.xlsx", (("same", "same"), ("left", "right")))
        source_preview = inspect_import_source(XlsxSource(path, sheet_name="Data"))
        preview = preview_import(
            source_preview,
            (ColumnMapping(1, "second"), ColumnMapping(0, "first")),
            (
                TargetColumn("first", ImportValueKind.TEXT, nullable=False),
                TargetColumn("second", ImportValueKind.TEXT, nullable=False),
            ),
            self._rules(),
        )
        self.assertEqual(tuple(column.label for column in source_preview.columns), ("same", "same"))
        self.assertEqual(preview.rows[0].values, ("right", "left"))

    def test_xlsx_empty_text_and_blank_cell_use_separate_null_rules(self) -> None:
        path = self._xlsx_file("empty-vs-blank.xlsx", (("empty_text", "blank_text"), ("", "")))
        self._replace_cell_with_shared_empty_text(path, "A2")
        inspected = inspect_import_source(XlsxSource(path, sheet_name="Data"))
        preview = preview_import(
            inspected,
            (ColumnMapping(0, "empty_text"), ColumnMapping(1, "blank_text")),
            (
                TargetColumn("empty_text", ImportValueKind.TEXT, nullable=True),
                TargetColumn("blank_text", ImportValueKind.TEXT, nullable=True),
            ),
            self._rules(),
        )
        self.assertEqual(preview.issues, ())
        self.assertEqual(preview.rows[0].values, ("", None))

    def test_formula_in_preview_and_beyond_preview_is_rejected(self) -> None:
        preview_formula = self._xlsx_file("formula-preview.xlsx", (("value",), ("=1+1",)))
        with self.assertRaises(ImportPreviewError) as raised:
            inspect_import_source(XlsxSource(preview_formula, sheet_name="Data"))
        self.assertEqual(raised.exception.code, ImportErrorCode.FORMULA_NOT_ALLOWED.value)

        later_formula = self._xlsx_file("formula-later.xlsx", (("value",), (1,), ("=1+1",)))
        limits = ImportLimits(max_preview_rows=1, max_chunk_rows=1)
        source = XlsxSource(later_formula, sheet_name="Data")
        inspected = inspect_import_source(source, limits=limits)
        self.assertEqual(len(inspected.rows), 1)
        self.assertTrue(inspected.has_more)
        preview = preview_import(
            inspected,
            (ColumnMapping(0, "value"),),
            (TargetColumn("value", ImportValueKind.INTEGER, nullable=False),),
            self._rules(),
        )
        approved = approve_import(preview, user_approved=True)
        chunks = iter_import_chunks(approved, chunk_rows=1)
        self.assertEqual(next(chunks).rows, ((1,),))
        with self.assertRaises(ImportPreviewError) as later:
            next(chunks)
        self.assertEqual(later.exception.code, ImportErrorCode.FORMULA_NOT_ALLOWED.value)

    def test_xlsx_extra_fields_and_zip_limits_fail_before_openpyxl(self) -> None:
        extra = self._xlsx_file("extra.xlsx", (("a", "b"), ("one", "two", "three")))
        with self.assertRaises(ImportPreviewError) as bad_row:
            inspect_import_source(XlsxSource(extra, sheet_name="Data"))
        self.assertEqual(bad_row.exception.code, ImportErrorCode.MALFORMED_ROW.value)

        valid = self._xlsx_file("limited.xlsx", (("a",), ("one",)))
        for limits in (
            ImportLimits(max_xlsx_entries=1),
            ImportLimits(max_xlsx_uncompressed_bytes=1),
            ImportLimits(max_rows=1),
        ):
            with self.subTest(limits=limits):
                with patch("openpyxl.load_workbook") as loader:
                    with self.assertRaises(ImportPreviewError) as limited:
                        inspect_import_source(
                            XlsxSource(valid, sheet_name="Data"),
                            limits=limits,
                        )
                self.assertEqual(limited.exception.code, ImportErrorCode.LIMIT_EXCEEDED.value)
                loader.assert_not_called()

    @staticmethod
    def _replace_xlsx_number(path: Path, cell_reference: str, lexical_value: str) -> None:
        temporary = path.with_suffix(".rewrite.xlsx")
        with zipfile.ZipFile(path, "r") as source_archive:
            entries = [(info.filename, source_archive.read(info.filename)) for info in source_archive.infolist()]
        sheet_name = next(name for name, _ in entries if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", name))
        updated = False
        rewritten = []
        pattern = re.compile(
            rf"(<c\b[^>]*\br=\"{re.escape(cell_reference)}\"[^>]*>.*?<v>)(.*?)(</v>.*?</c>)",
            re.DOTALL,
        )
        for name, content in entries:
            if name == sheet_name:
                text = content.decode("utf-8")
                text, count = pattern.subn(lambda match: match.group(1) + lexical_value + match.group(3), text, count=1)
                updated = count == 1
                content = text.encode("utf-8")
            rewritten.append((name, content))
        if not updated:
            raise AssertionError("expected numeric XML cell was not found")
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as destination:
            for name, content in rewritten:
                destination.writestr(name, content)
        os.replace(temporary, path)

    @staticmethod
    def _replace_cell_with_shared_empty_text(path: Path, cell_reference: str) -> None:
        temporary = path.with_suffix(".shared-empty.xlsx")
        with zipfile.ZipFile(path, "r") as source_archive:
            entries = [(info.filename, source_archive.read(info.filename)) for info in source_archive.infolist()]
        sheet_name = next(name for name, _ in entries if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", name))
        rewritten: list[tuple[str, bytes]] = []
        patched_cell = False
        patched_types = False
        patched_relationships = False
        cell_pattern = re.compile(rf"<c\b[^>]*\br=\"{re.escape(cell_reference)}\"[^>]*>.*?</c>")
        for name, content in entries:
            if name == sheet_name:
                text = content.decode("utf-8")
                text, count = cell_pattern.subn(
                    f'<c r="{cell_reference}" t="s"><v>0</v></c>',
                    text,
                    count=1,
                )
                patched_cell = count == 1
                content = text.encode("utf-8")
            elif name == "[Content_Types].xml":
                text = content.decode("utf-8")
                text = text.replace(
                    "</Types>",
                    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
                )
                patched_types = text != content.decode("utf-8")
                content = text.encode("utf-8")
            elif name == "xl/_rels/workbook.xml.rels":
                text = content.decode("utf-8")
                text = text.replace(
                    "</Relationships>",
                    '<Relationship Id="rIdSharedEmpty" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
                )
                patched_relationships = text != content.decode("utf-8")
                content = text.encode("utf-8")
            rewritten.append((name, content))
        if not patched_cell or not patched_types or not patched_relationships:
            raise AssertionError("expected XLSX parts were not found")
        rewritten.append(
            (
                "xl/sharedStrings.xml",
                b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1"><si><t></t></si></sst>',
            )
        )
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as destination:
            for name, content in rewritten:
                destination.writestr(name, content)
        os.replace(temporary, path)

    @staticmethod
    def _repack_xlsx_stored(path: Path) -> None:
        temporary = path.with_suffix(".stored.xlsx")
        with zipfile.ZipFile(path, "r") as source_archive:
            entries = [(info.filename, source_archive.read(info.filename)) for info in source_archive.infolist()]
        with zipfile.ZipFile(temporary, "w") as destination:
            for name, content in entries:
                info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_STORED
                destination.writestr(info, content)
        os.replace(temporary, path)


if __name__ == "__main__":
    unittest.main()
