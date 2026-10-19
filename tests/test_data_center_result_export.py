from __future__ import annotations

import csv
import datetime as dt
import json
import os
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest import mock

from tools.data_center.manual_session import ManualColumn, ManualQueryPage
from tools.data_center.result_export import (
    ExportError,
    ExportLimits,
    ExportScope,
    ExportSnapshot,
    ExportTarget,
    decode_csv_cell,
    decode_typed_cell,
    write_result_export,
)


class ResultExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="pengtools-result-export-")
        self.root = Path(self.temp_dir.name)
        self.target = ExportTarget("query-7", "tab-a", 4)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    @staticmethod
    def values():
        return (
            None,
            "",
            Decimal("12345678901234567890.00100"),
            9_007_199_254_740_993,
            b"\x00\xffPengTools",
            dt.datetime(2026, 10, 2, 9, 30, 0, 1200, tzinfo=dt.timezone(dt.timedelta(hours=5, minutes=30), "IST")),
            "=HYPERLINK(\"https://example.invalid\",\"x\")",
            "~PTX-CELL/1~literal-marker-text",
            dt.date(2020, 2, 29),
            -0.125,
        )

    def snapshot(self, rows=None, *, page_number=None, columns=None):
        values = self.values() if rows is None else rows
        names = tuple(f"c{i}" for i in range(len(values[0]))) if values else ("empty",)
        return ExportSnapshot(self.target, tuple(columns or names), tuple(tuple(row) for row in values), page_number)

    def test_csv_and_json_roundtrip_typed_values_duplicate_columns_and_marker_strings(self) -> None:
        values = self.values()
        snapshot = self.snapshot((values,), columns=("duplicate", "duplicate", *[f"c{i}" for i in range(2, len(values))]))

        csv_path = self.root / "typed.csv"
        write_result_export(csv_path, "csv", ExportScope.SELECTED, target=self.target, snapshot=snapshot)
        with csv_path.open("r", encoding="utf-8", newline="") as stream:
            records = list(csv.reader(stream))
        self.assertEqual(records[0], list(snapshot.columns))
        self.assertTrue(records[1][6].startswith("~PTX-CELL/1~"))
        decoded_csv = tuple(decode_csv_cell(value) for value in records[1])
        self.assertEqual(decoded_csv, values)

        json_path = self.root / "typed.json"
        write_result_export(json_path, "json", ExportScope.SELECTED, target=self.target, snapshot=snapshot)
        document = json.loads(json_path.read_text(encoding="utf-8"))
        self.assertEqual(document["contract"], "pengtools.typed-result-export")
        self.assertEqual(document["version"], 1)
        self.assertEqual(document["columns"], list(snapshot.columns))
        decoded_json = tuple(decode_typed_cell(cell) for cell in document["rows"][0])
        self.assertEqual(decoded_json, values)
        self.assertEqual(document["rows"][0][0], {"t": "null"})
        self.assertEqual(document["rows"][0][1], {"t": "str", "v": ""})

    def test_csv_escapes_formula_like_headers_and_values(self) -> None:
        path = self.root / "formula.csv"
        snapshot = ExportSnapshot(self.target, ("=header",), (("  +SUM(A1:A2)",),))
        write_result_export(path, "csv", ExportScope.SELECTED, target=self.target, snapshot=snapshot)
        with path.open("r", encoding="utf-8", newline="") as stream:
            records = list(csv.reader(stream))
        self.assertTrue(records[0][0].startswith("~PTX-CELL/1~"))
        self.assertTrue(records[1][0].startswith("~PTX-CELL/1~"))
        self.assertEqual(tuple(decode_csv_cell(row[0]) for row in records), ("=header", "  +SUM(A1:A2)"))

    def test_xlsx_roundtrip_metadata_formula_safety_and_long_text(self) -> None:
        import openpyxl

        long_text = "😀" * 20_000
        values = (
            None,
            "",
            Decimal("1.23000000000000000001"),
            10**30 + 7,
            b"\x00binary\xff",
            dt.datetime(2025, 11, 2, 1, 30, tzinfo=dt.timezone(dt.timedelta(hours=-4), "EDT")),
            "=1+1",
            long_text,
        )
        snapshot = self.snapshot((values,))
        path = self.root / "typed.xlsx"
        write_result_export(path, "xlsx", ExportScope.CURRENT_PAGE, target=self.target,
                            snapshot=ExportSnapshot(self.target, snapshot.columns, snapshot.rows, 1))

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            data = workbook["Data"]
            visible = list(data.iter_rows(values_only=False))[1]
            self.assertEqual(visible[6].value, "=1+1")
            self.assertEqual(visible[6].data_type, "s")
            self.assertTrue(visible[7].value.startswith("[full value"))
            meta = workbook["_PT_EXPORT_V1"]
            metadata_records = list(meta.iter_rows(values_only=True))
            self.assertEqual(metadata_records[0][2], "current_page")
            grouped: dict[tuple[int, int], list[tuple[int, int, str]]] = {}
            for record_index, column_index, payload, chunk_index, chunk_count in metadata_records[2:]:
                grouped.setdefault((record_index, column_index), []).append((chunk_index, chunk_count, payload))
            restored = []
            for col_index in range(len(values)):
                chunks = sorted(grouped[(0, col_index)])
                self.assertEqual(len(chunks), chunks[0][1])
                self.assertTrue(all(len(chunk.encode("utf-16-le")) // 2 <= 30_000 for _, _, chunk in chunks))
                encoded = "".join(chunk for _, _, chunk in chunks)
                restored.append(decode_typed_cell(json.loads(encoded)))
            self.assertEqual(tuple(restored), values)
            self.assertEqual(meta.sheet_state, "veryHidden")
        finally:
            workbook.close()

    def page(self, page_number, rows, *, has_more=False, target=None, truncated=False, columns=None):
        identity = target or self.target
        return ManualQueryPage(
            query_id=identity.query_id,
            tab_id=identity.tab_id,
            generation=identity.generation,
            columns=columns or (ManualColumn("v"),),
            rows=tuple(tuple(row) for row in rows),
            page_number=page_number,
            page_size=200,
            has_more=has_more,
            result_bytes=0,
            truncated_by_budget=truncated,
            rowcount=None,
        )

    def test_full_query_requires_complete_contiguous_untruncated_matching_chain(self) -> None:
        valid = [self.page(1, (("one",),), has_more=True), self.page(2, (("two",),))]
        path = self.root / "full.json"
        receipt = write_result_export(path, "json", ExportScope.FULL_QUERY, target=self.target, pages=iter(valid))
        self.assertEqual(receipt.rows, 2)
        self.assertEqual([decode_typed_cell(cell) for row in json.loads(path.read_text(encoding="utf-8"))["rows"] for cell in row], ["one", "two"])

        invalid_sources = (
            [self.page(1, (("secret",),), has_more=True)],
            [self.page(2, (("secret",),))],
            [self.page(1, (("secret",),), truncated=True)],
            [self.page(1, (("secret",),)), self.page(2, (("secret",),))],
            [self.page(1, (("secret",),), target=ExportTarget("other", "tab-a", 4))],
        )
        for index, source in enumerate(invalid_sources):
            target_path = self.root / f"invalid-{index}.json"
            with self.assertRaises(ExportError) as caught:
                write_result_export(target_path, "json", "full_query", target=self.target, pages=source)
            self.assertIn(caught.exception.code, {"INCOMPLETE_PAGE_CHAIN", "SOURCE_TRUNCATED", "TARGET_MISMATCH"})
            self.assertNotIn("secret", str(caught.exception))
            self.assertFalse(target_path.exists())

    def test_full_query_rejects_column_metadata_drift_but_preserves_duplicate_names(self) -> None:
        first = self.page(1, ((1,),), has_more=True, columns=(ManualColumn("v", "integer"),))
        changed = self.page(2, ((2,),), columns=(ManualColumn("v", "changed-type"),))
        path = self.root / "metadata-drift.json"
        with self.assertRaises(ExportError) as caught:
            write_result_export(path, "json", "full_query", target=self.target, pages=(first, changed))
        self.assertEqual(caught.exception.code, "COLUMN_METADATA_MISMATCH")
        self.assertFalse(path.exists())

        duplicate_columns = (ManualColumn("same", "integer"), ManualColumn("same", "integer"))
        duplicate_pages = (
            self.page(1, ((1, 2),), has_more=True, columns=duplicate_columns),
            self.page(2, ((3, 4),), columns=duplicate_columns),
        )
        duplicate_path = self.root / "duplicate-columns.json"
        receipt = write_result_export(duplicate_path, "json", "full_query", target=self.target,
                                      pages=duplicate_pages)
        self.assertEqual(receipt.rows, 2)
        self.assertEqual(json.loads(duplicate_path.read_text(encoding="utf-8"))["columns"], ["same", "same"])

    def test_cancel_removes_own_temp_and_preserves_neighbors(self) -> None:
        destination = self.root / "cancel.csv"
        other = self.root / "keep.txt"
        other.write_bytes(b"preserve")
        calls = 0

        def cancel_after_opening():
            nonlocal calls
            calls += 1
            return calls >= 4

        rows = tuple((f"row-{i}",) for i in range(10))
        snapshot = ExportSnapshot(self.target, ("value",), rows)
        with self.assertRaises(ExportError) as caught:
            write_result_export(destination, "csv", "selected", target=self.target, snapshot=snapshot,
                                cancel_check=cancel_after_opening)
        self.assertEqual(caught.exception.code, "CANCELLED")
        self.assertFalse(destination.exists())
        self.assertEqual(other.read_bytes(), b"preserve")
        self.assertEqual(list(self.root.iterdir()), [other])

    def test_no_overwrite_publish_is_race_safe_and_preserves_competing_file(self) -> None:
        destination = self.root / "race.json"
        snapshot = ExportSnapshot(self.target, ("x",), ((1,),))
        real_link = os.link
        calls = 0

        def create_competing_file_then_link(source, dest):
            nonlocal calls
            calls += 1
            if calls == 1:
                Path(dest).write_bytes(b"other writer")
            return real_link(source, dest)

        with mock.patch("tools.data_center.result_export.os.link", side_effect=create_competing_file_then_link):
            with self.assertRaises(ExportError) as caught:
                write_result_export(destination, "json", "selected", target=self.target, snapshot=snapshot)
        self.assertEqual(caught.exception.code, "DESTINATION_EXISTS")
        self.assertEqual(destination.read_bytes(), b"other writer")
        self.assertEqual([path.name for path in self.root.iterdir()], [destination.name])

    def test_explicit_overwrite_replaces_only_after_complete_export(self) -> None:
        destination = self.root / "replace.json"
        destination.write_bytes(b"previous")
        snapshot = ExportSnapshot(self.target, ("x",), ((1,),))
        with self.assertRaises(ExportError) as caught:
            write_result_export(destination, "json", "selected", target=self.target, snapshot=snapshot)
        self.assertEqual(caught.exception.code, "DESTINATION_EXISTS")
        self.assertEqual(destination.read_bytes(), b"previous")

        receipt = write_result_export(destination, "json", "selected", target=self.target,
                                      snapshot=snapshot, overwrite=True)
        self.assertEqual(receipt.rows, 1)
        self.assertNotEqual(destination.read_bytes(), b"previous")

    def test_output_and_snapshot_limits_fail_closed_without_partial_files(self) -> None:
        destination = self.root / "limited.csv"
        snapshot = ExportSnapshot(self.target, ("v",), (("x" * 100,),))
        with self.assertRaises(ExportError) as caught:
            write_result_export(destination, "csv", "selected", target=self.target, snapshot=snapshot,
                                limits=ExportLimits(max_output_bytes=40))
        self.assertEqual(caught.exception.code, "OUTPUT_LIMIT")
        self.assertFalse(destination.exists())
        self.assertEqual(list(self.root.iterdir()), [])

    def test_xlsx_generation_limit_does_not_publish_partial_workbook(self) -> None:
        from openpyxl.worksheet import _writer as openpyxl_writer

        destination = self.root / "limited.xlsx"
        snapshot = ExportSnapshot(self.target, ("v",), (("value",),))
        before = set(openpyxl_writer.ALL_TEMP_FILES)
        with self.assertRaises(ExportError) as caught:
            write_result_export(destination, "xlsx", "selected", target=self.target, snapshot=snapshot,
                                limits=ExportLimits(max_output_bytes=100))
        self.assertEqual(caught.exception.code, "OUTPUT_LIMIT")
        self.assertFalse(destination.exists())
        self.assertEqual(list(self.root.iterdir()), [])
        self.assertEqual(set(openpyxl_writer.ALL_TEMP_FILES), before)

    def test_xlsx_metadata_sheet_obeys_physical_row_cap_including_headers(self) -> None:
        from openpyxl.worksheet import _writer as openpyxl_writer

        destination = self.root / "metadata-row-cap.xlsx"
        snapshot = ExportSnapshot(self.target, ("v",), tuple((f"row-{i}",) for i in range(5)))
        before = set(openpyxl_writer.ALL_TEMP_FILES)
        # A tiny physical limit exercises the real append/cleanup path without
        # constructing a million metadata rows.
        with mock.patch("tools.data_center.result_export._XLSX_MAX_WORKSHEET_ROWS", 6, create=True):
            with self.assertRaises(ExportError) as caught:
                write_result_export(
                    destination,
                    "xlsx",
                    "selected",
                    target=self.target,
                    snapshot=snapshot,
                    limits=ExportLimits(max_xlsx_metadata_rows=100),
                )
        self.assertEqual(caught.exception.code, "XLSX_METADATA_LIMIT")
        self.assertFalse(destination.exists())
        self.assertEqual(list(self.root.iterdir()), [])
        self.assertEqual(set(openpyxl_writer.ALL_TEMP_FILES), before)

    def test_xlsx_cancellation_cleans_spools_and_does_not_publish(self) -> None:
        from openpyxl.worksheet import _writer as openpyxl_writer

        destination = self.root / "cancelled.xlsx"
        snapshot = ExportSnapshot(self.target, ("v",), (("first",), ("second",)))
        before = set(openpyxl_writer.ALL_TEMP_FILES)
        calls = 0

        def cancel_during_rows():
            nonlocal calls
            calls += 1
            return calls >= 6

        with self.assertRaises(ExportError) as caught:
            write_result_export(destination, "xlsx", "selected", target=self.target,
                                snapshot=snapshot, cancel_check=cancel_during_rows)
        self.assertEqual(caught.exception.code, "CANCELLED")
        self.assertFalse(destination.exists())
        self.assertEqual(list(self.root.iterdir()), [])
        self.assertEqual(set(openpyxl_writer.ALL_TEMP_FILES), before)

    def test_unsupported_values_and_oversized_cells_have_safe_codes(self) -> None:
        for value, limit, expected in ((object(), ExportLimits(), "UNSUPPORTED_VALUE"),
                                       ("x" * 8, ExportLimits(max_cell_bytes=4), "CELL_LIMIT")):
            path = self.root / "bad.json"
            with self.assertRaises(ExportError) as caught:
                write_result_export(path, "json", "selected", target=self.target,
                                    snapshot=ExportSnapshot(self.target, ("v",), ((value,),)), limits=limit)
            self.assertEqual(caught.exception.code, expected)
            self.assertFalse(path.exists())


if __name__ == "__main__":
    unittest.main()
